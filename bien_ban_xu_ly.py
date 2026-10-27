"""Biên bản xử lý mẫu Hatico — tích hợp Streamlit, không lưu/chỉnh sửa dữ liệu nguồn.

Chức năng chỉ điền các trường có dữ liệu hoặc đã được người dùng xác nhận.
Giữ nguyên phần đầu trang, quy trình xử lý và công thức của mẫu DOCX.
"""
from __future__ import annotations

from copy import deepcopy
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
import hashlib
import re
import unicodedata

import pandas as pd

FORM_FILE = 'Mau_Bien_Ban_Xu_Ly_Mau_Tu_Dong.docx'
WORD_COLUMNS = [
    'Ngày PT', 'Tên mẫu', 'Csurr trước', 'TG giải hấp', 'KH ống', 'Vial',
    'CIS', 'Tên chỉ tiêu', 'K', 'Csurr sau', 'H%', 'Pa',
    'Nhiệt độ', 'V khí', 'C đo', 'Kết quả',
]
SOURCE_COLS = ['Tên mẫu', 'Tên chỉ tiêu', 'Kết quả', 'Ô nguồn']
WORD_FIELDS = ['loai_mau', 'ma_me', 'nguoi_phan_tich', 'ngay_phan_tich',
               'phuong_phap', 'chi_tieu_phan_tich']
CHEM_FIELDS = ['Hóa chất', 'Ngày pha', 'Hạn sử dụng']
# Phân nhóm theo tiền tố chính thức trên mã mẫu, không dựa trên đơn vị từ Excel.
SAMPLE_FAMILIES = {
    'KT': ('Khí thải', 'mg/Nm³'),
    'KLV': ('Khí làm việc', 'Mg/m³'),  # Viết đúng theo yêu cầu; chữ M hoa cần xác nhận SI.
    'KXQ': ('Khí xung quanh', 'µg/Nm³'),
    'NS': ('Nước sạch', ''), 'NM': ('Nước mặt', ''), 'NT': ('Nước thải', ''),
}
BLANK_MARKERS = {'blank', 'blk', 'mau trang', 'trang'}
# Giá trị đặt trước theo nền mẫu, không phải giá trị đo từ thiết bị.
DEFAULT_GAS_LITERS = {'KT': '24', 'KLV': '4', 'KXQ': '4'}



def _as_text(value):
    if value is None:
        return ''
    try:
        if pd.isna(value):
            return ''
    except (ValueError, TypeError):
        pass
    if isinstance(value, (date, datetime)):
        return value.strftime('%d/%m/%Y')
    # Dữ liệu kết quả dạng chuỗi KPH (MDL=...) được giữ nguyên dấu và cách viết.
    return str(value).strip()


def _norm(value):
    s = unicodedata.normalize('NFD', _as_text(value).casefold()).replace('đ', 'd')
    return ''.join(x for x in s if unicodedata.category(x) != 'Mn')


def sample_family(code):
    """Nhận diện tiền tố đầy đủ: KLV/KXQ không bị hiểu là KT."""
    text = _as_text(code).upper()
    m = re.match(r'^(KLV|KXQ|KT|NS|NM|NT)(?=[.\-\s]|\d|$)', text)
    return m.group(1) if m else ''


def infer_sample_type(codes):
    kinds = {sample_family(code) for code in codes}
    return SAMPLE_FAMILIES[next(iter(kinds))][0] if len(kinds) == 1 and '' not in kinds else ''


def split_sample_groups(frame):
    """Giữ nguyên từng kết quả, không bao giờ gộp hai nền mẫu vào một Word."""
    if frame.empty:
        raise ValueError('Không có dòng kết quả để phân nhóm.')
    grouped = {}
    unrecognized = []
    for _, record in frame.iterrows():
        family = sample_family(record.get('Tên mẫu', ''))
        if not family:
            unrecognized.append(_as_text(record.get('Tên mẫu', '')))
        else:
            grouped.setdefault(family, []).append(record)
    if unrecognized:
        raise ValueError('Chưa nhận diện nền mẫu theo tiền tố: ' + ', '.join(dict.fromkeys(unrecognized)) + '. Không được tự trộn nền mẫu.')
    return {name: pd.DataFrame(records).reset_index(drop=True) for name, records in grouped.items()}


def _is_blank_code(code):
    """Chỉ coi mã ghi rõ BLANK/MẪU TRẮNG là mẫu trắng thật."""
    value = _norm(code).strip()
    code_prefix = sample_family(code)
    if code_prefix:
        value = re.sub(r'^[a-z]+[.\-\s]*', '', value)
    return value in BLANK_MARKERS


def number_tubes_and_vials(frame, *, insert_blank=False):
    """Số vật lý theo từng mẫu: Blank = 1, mẫu thật bắt đầu 2.

    Nhiều chỉ tiêu của cùng mã mẫu cùng KH ống/Vial. Đây là quy ước đánh số,
    không phải số đo lấy từ thiết bị. Không đưa kết quả đo giả vào dòng Blank.
    """
    if frame.empty:
        raise ValueError('Không có dòng mẫu để đánh số.')
    base = frame.copy().reset_index(drop=True)
    if base['Tên mẫu'].map(_as_text).eq('').any():
        raise ValueError('Có mã mẫu trống.')
    sample_keys = list(dict.fromkeys(base['Tên mẫu'].map(_as_text)))
    real_blank = [sample for sample in sample_keys if _is_blank_code(sample)]
    if len(real_blank) > 1:
        raise ValueError('Nhiều mã mẫu có nhãn BLANK trong một nền; kiểm tra trước khi đánh số.')
    if real_blank:
        positions = {sample: number for number, sample in enumerate([real_blank[0]] + [x for x in sample_keys if x != real_blank[0]], 1)}
        base = pd.concat([base.loc[base['Tên mẫu'].eq(real_blank[0])], base.loc[~base['Tên mẫu'].eq(real_blank[0])]], ignore_index=True)
    else:
        positions = {sample: number for number, sample in enumerate(sample_keys, 2)}
        if insert_blank:
            row = {key: '' for key in base.columns}
            row['Tên mẫu'] = 'BLANK'
            row['K'] = '1'
            row['Ô nguồn'] = 'DÒNG BLANK MẶC ĐỊNH — chưa có kết quả phân tích'
            row['Chọn'] = True
            row['_dong_blank_mac_dinh'] = True
            base = pd.concat([pd.DataFrame([row]), base], ignore_index=True)
            positions['BLANK'] = 1
    # Khóa thứ tự theo số mẫu, không phải theo số lượng dòng chỉ tiêu.
    base['KH ống'] = base['Tên mẫu'].map(positions).map(str)
    base['Vial'] = base['KH ống']
    base['K'] = base['K'].map(_as_text).replace('', '1')
    # Ngày giải hấp theo ngày PT. Bỏ trống nếu bản thân ngày PT chưa có trong nguồn.
    # BLANK thêm tự động không có ngày PT cũng như không có thể tích mẫu giả.
    for field in ('Ngày PT', 'TG giải hấp', 'V khí'):
        if field not in base:
            base[field] = ''
    base['_tg_mac_dinh'] = False
    base['_v_khi_mac_dinh'] = False
    for idx in base.index:
        if bool(base.at[idx, '_dong_blank_mac_dinh']) if '_dong_blank_mac_dinh' in base.columns and pd.notna(base.at[idx, '_dong_blank_mac_dinh']) else False:
            continue
        # Quy tắc bắt buộc: TG giải hấp luôn là ngày PT (không dùng giá trị khác).
        if _as_text(base.at[idx, 'TG giải hấp']) != _as_text(base.at[idx, 'Ngày PT']):
            base.at[idx, '_tg_mac_dinh'] = True
        base.at[idx, 'TG giải hấp'] = _as_text(base.at[idx, 'Ngày PT'])
        if not _is_blank_code(base.at[idx, 'Tên mẫu']) and not _as_text(base.at[idx, 'V khí']):
            vol = DEFAULT_GAS_LITERS.get(sample_family(base.at[idx, 'Tên mẫu']), '')
            if vol:
                base.at[idx, 'V khí'] = vol
                base.at[idx, '_v_khi_mac_dinh'] = True
    return base

def source_table(raw, mode):
    """Tách cột thông tin thao tác khỏi các cột kết quả nguồn (chỉ đọc)."""
    if not isinstance(raw, pd.DataFrame) or raw.empty:
        raise ValueError('Chưa có bảng kết quả hợp lệ.')
    for required in ('Tên mẫu', 'Tên chỉ tiêu', 'Kết quả'):
        if required not in raw:
            raise ValueError(f'Bảng nguồn thiếu cột {required}.')
    data = raw.copy()
    for col in WORD_COLUMNS + ['Ô nguồn', 'Đơn vị kết quả nguồn', 'Mã mẻ nguồn', 'Phương pháp phân tích nguồn']:
        if col not in data:
            data[col] = ''
    for col in WORD_COLUMNS + ['Ô nguồn']:
        data[col] = data[col].map(_as_text)
    missing = int(data['Kết quả'].eq('').sum())
    data = data.loc[data['Kết quả'].ne('')].copy()
    if data.empty:
        raise ValueError('Không có kết quả thực để điền. Các ô trống không được suy đoán là KPH hoặc 0.')
    # Nguồn Excel không có các giá trị này. Bộ BB-KT cũ dùng mặc định số Vial,
    # V khí và K; không được coi những giá trị suy đoán đó là số đo thực tế.
    if mode == 'excel':
        for col in ('TG giải hấp', 'KH ống', 'Vial', 'V khí'):
            data[col] = ''
    else:
        data.loc[data['TG giải hấp'].eq(data['Ngày PT']), 'TG giải hấp'] = ''
        data['KH ống'] = ''
        data['Vial'] = ''
    data['K'] = data['K'].map(_as_text).replace('', '1')
    keys = set()
    for _, row in data.iterrows():
        if not _as_text(row['Tên mẫu']) or not _as_text(row['Tên chỉ tiêu']):
            raise ValueError('Có dòng kết quả thiếu mã mẫu hoặc chỉ tiêu.')
        key = (_norm(row['Tên mẫu']), _norm(row['Tên chỉ tiêu']))
        if key in keys:
            raise ValueError('Trùng mã mẫu/chỉ tiêu: ' + ' / '.join(key))
        keys.add(key)
    if 'Chọn' not in data:
        data.insert(0, 'Chọn', True)
    else:
        data['Chọn'] = True
    return data.reset_index(drop=True), missing


def suggest_metadata(frame, mode):
    codes = frame['Tên mẫu'].map(_as_text).tolist()
    def unique(field):
        if field not in frame:
            return ''
        vals = list(dict.fromkeys(x for x in frame[field].map(_as_text) if x))
        return vals[0] if len(vals) == 1 else ''
    batch = unique('Mã mẻ nguồn')
    date_found = unique('Ngày PT')
    source_method = unique('Phương pháp phân tích nguồn') if mode == 'excel' else ''
    units = unique('Đơn vị kết quả nguồn')
    indicators = list(dict.fromkeys(_as_text(x) for x in frame['Tên chỉ tiêu'] if _as_text(x)))
    all_indicators = '; '.join(indicators)
    # Tránh ép phần SOP của mẫu Word sang trang khác vì tên nhóm chỉ tiêu quá dài.
    header_indicators = (f'{len(indicators)} chỉ tiêu (chi tiết tại bảng kết quả)'
                         if len(all_indicators) > 150 else all_indicators)
    return {'ma_me': batch, 'ngay_phan_tich': date_found,
            'loai_mau': infer_sample_type(codes), 'phuong_phap_nguon': source_method,
            'don_vi_goi_y': units, 'chi_tieu_phan_tich': header_indicators,
            'danh_sach_chi_tieu': all_indicators}


def _replace_paragraph_tokens(paragraph, mapping):
    """Thay marker nhưng giữ kiểu chữ và các run tĩnh của mẫu Word.

    Word có thể chia một {{ marker }} thành nhiều run. Chỉ phần run nằm
    trong marker được cập nhật, không gom cả paragraph vào một run.
    """
    if '{{' not in paragraph.text:
        return
    runs = paragraph.runs
    if not runs:
        return
    text = ''.join(run.text for run in runs)
    pattern = re.compile(r'\{\{\s*([a-zA-Z0-9_]+)\s*\}\}')
    ranges = []
    cursor = 0
    for i, run in enumerate(runs):
        ranges.append((cursor, cursor + len(run.text), i))
        cursor += len(run.text)
    # Đi từ phải sang trái để vị trí gốc của các marker trước không đổi.
    for match in reversed(list(pattern.finditer(text))):
        if match.group(1) not in mapping:
            continue
        first = next((i for a, b, i in ranges if a <= match.start() < b), None)
        last = next((i for a, b, i in ranges if a < match.end() <= b), None)
        if first is None or last is None:
            continue
        a = ranges[first][0]
        b = ranges[last][0]
        replacement = _as_text(mapping[match.group(1)])
        if first == last:
            original = runs[first].text
            runs[first].text = original[:match.start()-a] + replacement + original[match.end()-a:]
        else:
            # Trả phần suffix về đúng run cuối, giữ nguyên định dạng text ngoài marker.
            suffix = runs[last].text[match.end()-b:]
            runs[first].text = runs[first].text[:match.start()-a] + replacement
            for i in range(first + 1, last):
                runs[i].text = ''
            runs[last].text = suffix


def _all_paragraphs(document):
    for p in document.paragraphs:
        yield p
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p
    for section in document.sections:
        for p in section.header.paragraphs:
            yield p
        for p in section.footer.paragraphs:
            yield p


def _get_table(document, tokens, min_cols=1):
    for table in document.tables:
        if len(table.columns) < min_cols:
            continue
        content = ' '.join(cell.text for row in table.rows[:2] for cell in row.cells)
        if all(token in content for token in tokens):
            return table
    raise ValueError('Mẫu Word không có bảng phù hợp: ' + ', '.join(tokens))


def _set_cell(cell, value):
    val = _as_text(value)
    if not cell.paragraphs:
        cell.text = val
        return
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = val
        for run in p.runs[1:]:
            run.text = ''
    else:
        p.add_run(val)
    # Giữ nguyên định dạng paragraph và cell gốc, làm sạch các đoạn dư.
    for other in cell.paragraphs[1:]:
        for run in other.runs:
            run.text = ''


def _fill_table(document, rows, unit):
    """Bảng Word khổ ngang 16 cột: kích thước ổn định, hàng xen kẽ, tiêu đề lặp."""
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import (WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT,
                                 WD_ROW_HEIGHT_RULE)
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    table = _get_table(document, ('Ngày PT', 'KH mẫu', 'Csurr trước'), min_cols=16)
    if len(table.columns) != 16:
        raise ValueError('Bảng kết quả phải có đúng 16 cột theo mẫu Word đã cung cấp.')
    original_capacity = len(table.rows) - 1
    if original_capacity < 1:
        raise ValueError('Mẫu Word không có hàng dữ liệu dự phòng.')
    prototype = deepcopy(table.rows[-1]._tr)
    for _ in range(max(0, len(rows) - original_capacity)):
        table._tbl.append(deepcopy(prototype))
    for leftover in list(table.rows)[len(rows)+1:]:
        table._tbl.remove(leftover._tr)

    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Tổng rộng 26,55 cm, không vượt khổ ngang 27,94 cm trừ lề 0,51 cm mỗi bên.
    widths_cm = [1.75, 2.35, 1.50, 1.70, 1.10, 1.05, 1.05, 3.40,
                 1.00, 1.50, 1.05, 1.50, 1.55, 1.35, 1.70, 3.00]
    for i, width in enumerate(widths_cm):
        table.columns[i].width = Cm(width)

    def style_cell(cell, *, header=False, alternate=False, blank=False, col=0):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tcPr = cell._tc.get_or_add_tcPr()
        for node in tcPr.findall(qn('w:shd')):
            tcPr.remove(node)
        shade = OxmlElement('w:shd')
        shade.set(qn('w:fill'), 'DCEAF2' if header else ('EAF2F8' if blank else 'F4F7FB' if alternate else 'FFFFFF'))
        tcPr.append(shade)
        margins = tcPr.find(qn('w:tcMar'))
        if margins is None:
            margins = OxmlElement('w:tcMar'); tcPr.append(margins)
        for name, amount in [('top',80),('bottom',80),('left',70),('right',65)]:
            edge=margins.find(qn('w:'+name))
            if edge is None: edge=OxmlElement('w:'+name); margins.append(edge)
            edge.set(qn('w:w'), str(amount));edge.set(qn('w:type'),'dxa')
        borders=tcPr.find(qn('w:tcBorders'))
        if borders is None: borders=OxmlElement('w:tcBorders');tcPr.append(borders)
        for side in ('top','bottom','left','right'):
            edge=borders.find(qn('w:'+side))
            if edge is None:edge=OxmlElement('w:'+side);borders.append(edge)
            edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4')
            edge.set(qn('w:color'),'BACAD6')
        for para in cell.paragraphs:
            para.alignment = (WD_ALIGN_PARAGRAPH.LEFT if (col in (1,7,15) and not header)
                              else WD_ALIGN_PARAGRAPH.CENTER)
            fmt=para.paragraph_format
            fmt.space_before=Pt(0);fmt.space_after=Pt(0)
            fmt.line_spacing=1.05
            for run in para.runs:
                run.font.name='Times New Roman'
                run.font.size=Pt(7.4 if header else 7.5 if col==15 else 8.0)
                run.font.bold=header or blank
                if header:run.font.color.rgb=__import__('docx').shared.RGBColor(23,49,77)

    label = f'Kết quả\n({_as_text(unit)})' if _as_text(unit) else 'Kết quả\n(CHƯA XÁC NHẬN ĐƠN VỊ)'
    _set_cell(table.rows[0].cells[15], label)
    header_pr=table.rows[0]._tr.get_or_add_trPr()
    if header_pr.find(qn('w:tblHeader')) is None:
        header_pr.append(OxmlElement('w:tblHeader'))
    table.rows[0].height=Pt(57)
    table.rows[0].height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
    for ci, width in enumerate(widths_cm):
        cell=table.rows[0].cells[ci]; cell.width=Cm(width)
        style_cell(cell,header=True,col=ci)
    for i, (_, record) in enumerate(rows.iterrows(), 1):
        current=table.rows[i]
        current.height=Pt(26)
        current.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
        tr_pr=current._tr.get_or_add_trPr()
        if tr_pr.find(qn('w:cantSplit')) is None:tr_pr.append(OxmlElement('w:cantSplit'))
        is_blank=_is_blank_code(record.get('Tên mẫu',''))
        for ci,key in enumerate(WORD_COLUMNS):
            cell=current.cells[ci]
            cell.width=Cm(widths_cm[ci])
            _set_cell(cell, record.get(key,''))
            style_cell(cell,alternate=(i%2==0),blank=is_blank,col=ci)
    return max(original_capacity,len(rows))


def generate_bien_ban(template_bytes, metadata, chemicals, rows, result_unit, *, draft=False):
    """Xuất Word. Không đưa ra tính toán, không sửa kết quả nguyên bản."""
    from docx import Document
    if not template_bytes:
        raise ValueError('Thiếu file mẫu Word.')
    if rows.empty:
        raise ValueError('Chưa chọn dòng kết quả để xuất.')
    data = rows.copy()
    for col in WORD_COLUMNS:
        if col not in data:
            data[col] = ''
    document = Document(BytesIO(template_bytes))
    _get_table(document, ('{{ loai_mau }}', '{{ ma_me }}'))
    _get_table(document, ('{{ hc1_ten }}',))
    mapping = {key: _as_text(metadata.get(key, '')) for key in WORD_FIELDS}
    chemicals = list(chemicals)
    for i in range(4):
        entry = chemicals[i] if i < len(chemicals) else {}
        for token, field in [('ten', 'Hóa chất'), ('ngay_pha', 'Ngày pha'), ('han_su_dung', 'Hạn sử dụng')]:
            mapping[f'hc{i+1}_{token}'] = _as_text(entry.get(field, ''))
    for p in _all_paragraphs(document):
        _replace_paragraph_tokens(p, mapping)
    _fill_table(document, data, result_unit)
    if draft:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt
        results_table = _get_table(document, ('Ngày PT', 'KH mẫu', 'Csurr trước'), min_cols=16)
        p = OxmlElement('w:p')
        r = OxmlElement('w:r')
        r_pr = OxmlElement('w:rPr')
        b = OxmlElement('w:b'); r_pr.append(b)
        sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '18'); r_pr.append(sz)
        r.append(r_pr)
        t = OxmlElement('w:t')
        t.text = ('BẢN NHÁP – CẦN XÁC NHẬN ĐƠN VỊ, SOP VÀ DỮ LIỆU CÒN THIẾU. '
                  'P được sao chép nguyên giá trị từ Excel, chưa xác nhận đơn vị áp suất.')
        r.append(t); p.append(r)
        results_table._tbl.addprevious(p)
    unresolved = set()
    for p in _all_paragraphs(document):
        unresolved.update(re.findall(r'\{\{[^{}]+\}\}', p.text))
    if unresolved:
        raise ValueError('Mẫu Word còn thẻ chưa nhận diện: ' + ', '.join(sorted(unresolved)[:8]))
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _confirmed_chemicals(chem_editor):
    result = []
    for _, item in chem_editor.iterrows():
        result.append({key: _as_text(item.get(key, '')) for key in CHEM_FIELDS})
    return result


def _stock_options(stock):
    if not isinstance(stock, pd.DataFrame) or stock.empty or 'Tên Hóa Chất' not in stock:
        return []
    return sorted({str(v).strip() for v in stock['Tên Hóa Chất'].dropna() if str(v).strip()})


def _chem_rows(selected, stock):
    chemicals = []
    for name in selected[:4]:
        match = stock[stock['Tên Hóa Chất'].astype(str).eq(name)] if isinstance(stock,pd.DataFrame) and 'Tên Hóa Chất' in stock else pd.DataFrame()
        expiry = ''
        if len(match) == 1:
            expiry = _as_text(match.iloc[0].get('Hạn Sử Dụng', ''))
        chemicals.append({'Hóa chất': name, 'Ngày pha': '', 'Hạn sử dụng': expiry})
    while len(chemicals) < 4:
        chemicals.append({x: '' for x in CHEM_FIELDS})
    return pd.DataFrame(chemicals)


def _warning_rows(frame):
    warnings = []
    for field in ('C đo', 'Csurr trước', 'Csurr sau', 'H%', 'V khí', 'TG giải hấp', 'K'):
        if field not in frame:
            continue
        n = int(frame[field].map(_as_text).eq('').sum())
        if n:
            warnings.append(f'{field}: còn trống {n}/{len(frame)} dòng')
    return warnings


def generate_bien_ban_excel(metadata, rows, result_unit, *, draft=False):
    """Xuất Excel 16 cột từ cùng bảng đang đưa vào Word bằng artifact_tool.

    Giữ chuỗi kết quả gốc (kể cả KPH/MDL) và thêm sheet truy nguyên nguồn.
    Thực thi khi nhấn nút, không tính lại bất cứ chỉ tiêu nào.
    """
    from artifact_tool import Workbook, SpreadsheetFile
    wb = Workbook.create()
    ws = wb.worksheets.add('Biên bản')
    ws.merge_cells('A1:P2')
    ws.get_range('A1').values = [['BIÊN BẢN XỬ LÝ MẪU — ' + _as_text(metadata.get('loai_mau'))]]
    ws.get_range('A1:P2').format = {
        'fill': '#17324D', 'font': {'name': 'Arial', 'bold': True, 'size': 16, 'color': '#FFFFFF'},
        'horizontal_alignment': 'center', 'vertical_alignment': 'center', 'row_height': 28}
    ws.merge_cells('A3:H3');ws.merge_cells('I3:P3')
    ws.get_range('A3').values = [['Mã mẻ: '+_as_text(metadata.get('ma_me'))]]
    ws.get_range('I3').values = [['Ngày PT: '+_as_text(metadata.get('ngay_phan_tich'))]]
    ws.merge_cells('A4:H4');ws.merge_cells('I4:P4')
    ws.get_range('A4').values = [['Người phân tích: '+_as_text(metadata.get('nguoi_phan_tich'))]]
    ws.get_range('I4').values = [['Đơn vị kết quả: '+_as_text(result_unit)]]
    ws.merge_cells('A5:P5')
    ws.get_range('A5').values = [['Phương pháp xử lý mẫu: '+_as_text(metadata.get('phuong_phap'))]]
    ws.merge_cells('A6:P6')
    ws.get_range('A6').values = [['BẢN NHÁP — cần xác nhận dữ liệu/đơn vị/SOP trước khi phát hành.' if draft
                                else 'Dữ liệu tự động trích xuất; cần đối chiếu hồ sơ trước khi phát hành.']]
    ws.get_range('A3:P6').format = {
        'fill': '#F0F5F9','font':{'name':'Arial','size':10,'color':'#33485C'},
        'row_height':22,'vertical_alignment':'center'}
    ws.get_range('A6:P6').format = {'fill':'#FFF4DF','font':{'name':'Arial','bold':True,'size':10,'color':'#85591F'},'row_height':24}
    columns = ['Ngày PT','KH mẫu','Csurr trước','TG giải hấp','KH ống đựng mẫu',
               'Ký hiệu Vial','CIS','Chỉ tiêu','Hệ số pha loãng (K)',
               'Csurr sau','H%','Pa','Nhiệt độ (°C)','V khí (L)','C đo',
               'Kết quả ('+_as_text(result_unit)+')']
    ws.get_range('A8:P8').values = [columns]
    ws.get_range('A8:P8').format = {
        'fill':'#22516D','font':{'name':'Arial','bold':True,'size':10,'color':'#FFFFFF'},
        'horizontal_alignment':'center','vertical_alignment':'center',
        'wrap_text':True,'row_height':51}
    matrix=[]
    for _,row in rows.iterrows():
        values=[]
        for key in WORD_COLUMNS:
            val=_as_text(row.get(key,''))
            if key in ('KH ống','Vial','K','V khí','Pa','Nhiệt độ') and re.fullmatch(r'\d+(?:[.,]\d+)?',val):
                val = float(val.replace(',','.')) if ('.' in val or ',' in val) else int(val)
            values.append(val)
        matrix.append(values)
    if matrix:
        last=8+len(matrix)
        ws.get_range(f'A9:P{last}').values = matrix
        ws.get_range(f'A9:P{last}').format = {
            'font':{'name':'Arial','size':10,'color':'#203448'},
            'vertical_alignment':'center','horizontal_alignment':'center',
            'row_height':28, 'wrap_text':True}
        for i,entry in enumerate(matrix,9):
            if entry[1]=='BLANK':
                ws.get_range(f'A{i}:P{i}').format.fill='#DDECF7'
            elif i%2==0:
                ws.get_range(f'A{i}:P{i}').format.fill='#F5F8FB'
        for col in ('B','H','P'):
            ws.get_range(f'{col}9:{col}{last}').format.horizontal_alignment='left'
    for letter,width in zip('ABCDEFGHIJKLMNOP', [13,20,14,15,14,12,11,24,14,14,10,11,15,12,14,24]):
        ws.get_range(f'{letter}:{letter}').format.column_width=width
    ws.freeze_panes.freeze_rows(8)
    trace=wb.worksheets.add('Nguồn dữ liệu')
    source_columns=['KH mẫu','Chỉ tiêu','Kết quả nguyên gốc','Ô nguồn Excel',
                    'Ngày PT','TG giải hấp','V khí (L)','Pa (giá trị gốc)','Nhiệt độ (°C)','Lưu ý']
    trace.get_range('A1:J1').values=[source_columns]
    trace.get_range('A1:J1').format={'fill':'#22516D','font':{'bold':True,'color':'#FFFFFF','size':10},'row_height':34,'wrap_text':True}
    source_rows=[]
    for _,r in rows.iterrows():
        source_rows.append([_as_text(r.get('Tên mẫu')),_as_text(r.get('Tên chỉ tiêu')),_as_text(r.get('Kết quả')),
                            _as_text(r.get('Ô nguồn')),_as_text(r.get('Ngày PT')),
                            _as_text(r.get('TG giải hấp')),_as_text(r.get('V khí')),
                            _as_text(r.get('Pa')),_as_text(r.get('Nhiệt độ')),
                            'BLANK mặc định — không có số đo' if _as_text(r.get('Tên mẫu')).upper()=='BLANK' else 'V khí/TG giải hấp theo quy tắc mặc định nếu không có số liệu nguồn'])
    if source_rows:
        trace.get_range(f'A2:J{len(source_rows)+1}').values=source_rows
        trace.get_range(f'A2:J{len(source_rows)+1}').format={'font':{'size':10},'row_height':23,'wrap_text':True}
    for col,width in zip('ABCDEFGHIJ',[18,24,21,20,15,16,14,19,20,70]):
        trace.get_range(f'{col}:{col}').format.column_width=width
    trace.freeze_panes.freeze_rows(1)
    return SpreadsheetFile.export_xlsx(wb).data

def render_bien_ban_xu_ly(*, excel_reader, result_reader, results, stock=None, excel_exporter=None):
    """Widget riêng, dùng các hàm đọc Excel/phiên đã được kiểm tra trong app gốc."""
    import streamlit as st

    st.markdown('#### Tự động điền Biên bản xử lý mẫu (Word và Excel theo nền)')
    st.caption('Điền trường có căn cứ từ Excel hoặc phiên phân tích. Phần đầu trang, quy trình và công thức SOP trong Word giữ nguyên. Không ghi ngược kết quả vào Sheets.')

    template_up = st.file_uploader('Mẫu Word (.docx) — có thể để trống khi đã đặt file mẫu cạnh app.py', type=['docx'], key='bbxl_template')
    template_path = Path(__file__).resolve().with_name(FORM_FILE)
    template_bytes = template_up.getvalue() if template_up else template_path.read_bytes() if template_path.is_file() else None
    if not template_bytes:
        st.warning(f'Chưa có mẫu Word. Tải lên mẫu .docx hoặc đặt {FORM_FILE} cùng thư mục app.py.')

    source = st.radio('Nguồn dữ liệu', ['File kết quả mẻ Excel', 'Kết quả đã tính trong phiên'], horizontal=True, key='bbxl_source')
    is_excel = source.startswith('File')
    xlsx_up = st.file_uploader('Tải file KetQuaMeThuNghiem.xlsx', type=['xlsx'], key='bbxl_xlsx') if is_excel else None
    if is_excel and xlsx_up is None:
        st.info('Tải Excel kết quả mẻ để tự đọc mã mẻ, chỉ tiêu, kết quả, P và T; mỗi nền xuất riêng cả Word và Excel.')
        return
    if not is_excel and (not isinstance(results, pd.DataFrame) or results.empty):
        st.info('Chưa có kết quả phiên. Vào Tiện ích Phân tích và lưu kết quả trước.')
        return
    if not is_excel and st.session_state.get('results_stale'):
        st.error('Kết quả phiên đã cũ vì dữ liệu liên quan thay đổi. Hãy tính lại trước khi điền biên bản.')
        return

    source_bytes = xlsx_up.getvalue() if is_excel else None
    source_digest = hashlib.sha256(source_bytes if is_excel else results.to_json(orient='split', force_ascii=False).encode()).hexdigest()[:14]
    try:
        if is_excel:
            input_frame, parse_issues = excel_reader(source_bytes)
        else:
            input_frame, parse_issues = result_reader(results), []
        base_frame, ignored = source_table(input_frame, 'excel' if is_excel else 'session')
    except (ValueError, KeyError, TypeError) as exc:
        st.error('Không đọc được kết quả: ' + str(exc))
        return
    except Exception as exc:
        st.error(f'Lỗi tải dữ liệu: {exc}')
        return

    if ignored:
        st.info(f'Đã bỏ qua {ignored} dòng mẫu/chỉ tiêu chưa có kết quả; không tự coi ô trống là KPH.')
    important = [msg for msg in parse_issues if ('ô xám' in msg or 'Nhiều giá trị' in msg or 'công thức' in msg)]
    if important:
        with st.expander(f'Cảnh báo trích xuất ({len(important)})'):
            for msg in important[:30]: st.warning(msg)
    try:
        # Bản xem trước không thêm dòng BLANK; mẫu thực bắt đầu ở vị trí 2.
        base_frame = pd.concat([number_tubes_and_vials(group, insert_blank=False)
                                for group in split_sample_groups(base_frame).values()], ignore_index=True)
    except ValueError as exc:
        st.error(str(exc))
        return
    hints = suggest_metadata(base_frame, 'excel' if is_excel else 'session')
    st.success(f'Đọc được {len(base_frame)} dòng kết quả từ {base_frame["Tên mẫu"].nunique()} mẫu. Dữ liệu số không bị tính lại.')

    col1, col2, col3 = st.columns(3)
    with col1:
        batch = st.text_input('Mã mẻ', value=hints['ma_me'], key='bbxl_batch_'+source_digest)
    with col2:
        analyst = st.text_input('Người phân tích — xác nhận thủ công', key='bbxl_analyst_'+source_digest)
        analysis_date = st.text_input('Ngày phân tích', value=hints['ngay_phan_tich'], key='bbxl_date_'+source_digest)
    with col3:
        method = st.text_input('Phương pháp XỬ LÝ mẫu — xác nhận SOP', key='bbxl_method_'+source_digest)
    if hints['phuong_phap_nguon']:
        st.caption('Phương pháp PHÂN TÍCH ghi trong Excel: ' + hints['phuong_phap_nguon'] + '. Cần xác nhận SOP xử lý riêng cho mỗi nền.')
    families = list(split_sample_groups(base_frame))
    st.markdown('**Tách biên bản theo nền mẫu — Word và Excel**')
    st.caption('Mỗi nền: vị trí 1 = BLANK, mẫu thật 2, 3…; KH ống/Vial đánh theo mẫu; K = 1; TG giải hấp = Ngày PT; V khí KT = 24 L, KLV/KXQ = 4 L.')
    configured_units = {}
    for family in families:
        kind, default_unit = SAMPLE_FAMILIES[family]
        configured_units[family] = st.text_input(
            f'{family} — {kind}: đơn vị kết quả', value=default_unit,
            key='bbxl_unit_'+family+'_'+source_digest)
        if family == 'KLV':
            st.warning('KLV đang ghi **Mg/m³** theo đúng yêu cầu. Trong ký hiệu SI, Mg (megagam) khác mg (miligam). Cần xác nhận chữ M hoa trước khi phát hành.')
        if family in ('NS', 'NM', 'NT'):
            st.warning(f'{family}: chưa có đơn vị mặc định và mẫu Word đang mô tả quy trình xử lý khí. Phải chọn đúng SOP và đơn vị.')
    if base_frame['Pa'].map(_as_text).ne('').any():
        st.caption('Cột Pa lấy nguyên số từ P trong ghi chú Excel, không tự chuyển đơn vị áp suất vì nguồn có thể không ghi đơn vị.')
    if is_excel and base_frame['Nhiệt độ'].map(_as_text).ne('').any():
        st.caption('Đã đọc nhiệt độ từ ghi chú Excel (bao gồm T=33, t=30,14 hoặc t33).')

    st.markdown('**Hóa chất thực sự đã sử dụng**')
    options = _stock_options(stock)
    selected = st.multiselect('Gợi ý từ kho (không chứng minh đã sử dụng)', options, max_selections=4, key='bbxl_stockpick_'+source_digest) if options else []
    init_chem = _chem_rows(selected, stock)
    chem = st.data_editor(init_chem, num_rows='fixed', hide_index=True, use_container_width=True, key='bbxl_chems_'+source_digest+'_'+hashlib.sha1('|'.join(selected).encode()).hexdigest()[:7])
    st.caption('Ngày pha không có trong danh sách kho thì để trống cho đến khi đối chiếu sổ pha. Hạn dùng phải kiểm tra theo chai/lô thực tế.')

    st.markdown('**Bảng 16 cột — xem trước và bổ sung thông tin thao tác**')
    st.caption('K mặc định = 1; ống/Vial tự đánh số với BLANK ở vị trí 1; TG giải hấp = Ngày PT; V khí KT = 24 L, KLV/KXQ = 4 L (được sửa khi có thể tích thực). C đo không tự tạo.')
    preview_cols = ['Chọn'] + WORD_COLUMNS + ['Ô nguồn']
    source_disabled = [name for name in SOURCE_COLS + ['KH ống', 'Vial', 'TG giải hấp'] if name in preview_cols]
    if 'Chọn' in source_disabled: source_disabled.remove('Chọn')
    # Cho nhập các giá trị còn thiếu ở vùng thao tác; giữ nguyên các cột số liệu phân tích.
    editable = st.data_editor(base_frame[preview_cols], disabled=source_disabled,
                              hide_index=True, num_rows='fixed', use_container_width=True,
                              height=380, key='bbxl_rows_'+source_digest)
    # Bất biến dữ liệu: chỉ được điền giá trị đang trống; không sửa giá trị đọc từ nguồn.
    source_change = []
    manually_added = []
    for field in WORD_COLUMNS:
        for idx in editable.index:
            before = _as_text(base_frame.at[idx, field])
            after = _as_text(editable.at[idx, field])
            if before and before != after and field != 'K' and not (field == 'V khí' and bool(base_frame.at[idx, '_v_khi_mac_dinh'])) and not (field == 'TG giải hấp' and bool(base_frame.at[idx, '_tg_mac_dinh'])):
                source_change.append(f'{field}: mẫu {base_frame.at[idx, "Tên mẫu"]}')
            elif (not before and after) or (field == 'K' and before != after) or (field in ('V khí','TG giải hấp') and before != after):
                manually_added.append(f'{field}: mẫu {base_frame.at[idx, "Tên mẫu"]}')
    chosen = editable.loc[editable['Chọn'].fillna(False).astype(bool)].copy()
    if source_change:
        st.error('Không được chỉnh giá trị đã có từ nguồn: ' + '; '.join(source_change[:6]) + '. Khôi phục ô hoặc cập nhật tại nguồn.')
    if manually_added:
        st.info(f'Đã bổ sung thủ công {len(manually_added)} ô thiếu dữ liệu; cần đối chiếu sổ phân tích trước khi phát hành.')
    st.caption(f'Đã chọn {len(chosen)}/{len(base_frame)} dòng kết quả.')
    blanks = _warning_rows(chosen)
    if blanks:
        with st.expander('Dữ liệu thao tác còn thiếu, cần đối chiếu hồ sơ gốc', expanded=False):
            st.write('; '.join(blanks))

    draft = st.checkbox('Xuất BẢN NHÁP để rà soát', value=True, key='bbxl_draft_'+source_digest)
    confirm = st.checkbox('Tôi đã đối chiếu mã mẫu, chỉ tiêu và kết quả với Excel/hồ sơ gốc.', key='bbxl_confirm_'+source_digest)
    klv_confirm = st.checkbox('Tôi xác nhận chính xác đơn vị KLV viết hoa Mg/m³ (không phải mg/m³)',
                              key='bbxl_klv_capital_'+source_digest) if 'KLV' in families else True
    chem_rows = _confirmed_chemicals(chem)
    # Một đầu vào dẫn đến nhiều file. Kết quả ZIP để dễ tải đồng loạt và file riêng để đối chiếu.
    payload = (repr({'batch': batch, 'analyst': analyst, 'date': analysis_date, 'method': method,
                     'units': configured_units}) + repr(chem_rows) + chosen[WORD_COLUMNS].to_json(orient='records', force_ascii=False)
               + str(draft) + str(klv_confirm) + hashlib.sha256(template_bytes or b'').hexdigest())
    form_digest = hashlib.sha256(payload.encode()).hexdigest()[:16]

    if st.button('Xuất Word + Excel theo từng nền mẫu', type='primary',
                 disabled=not bool(template_bytes) or not confirm or chosen.empty or bool(source_change),
                 key='bbxl_generate_'+source_digest):
        try:
            from zipfile import ZipFile, ZIP_DEFLATED
            for value in chosen['K']:
                if not re.fullmatch(r'\d+(?:[.,]\d+)?', _as_text(value)) or float(str(value).replace(',', '.')) < 1:
                    raise ValueError('Hệ số pha loãng K phải là số thực ≥ 1.')
            if not draft and 'KLV' in families and configured_units['KLV'].strip()=='Mg/m³' and not klv_confirm:
                raise ValueError('Phải xác nhận cách viết Mg/m³ hoặc đổi sang mg/m³ trước khi xuất bản hoàn chỉnh.')
            if not draft and (not analyst.strip() or not method.strip() or not analysis_date.strip() or not batch.strip()):
                raise ValueError('Biên bản hoàn chỉnh yêu cầu mã mẻ, người phân tích, ngày và SOP xử lý mẫu.')
            if not draft and not any(_as_text(c['Hóa chất']) for c in chem_rows):
                raise ValueError('Biên bản hoàn chỉnh: phải xác nhận hóa chất thực sự đã sử dụng.')
            created = []
            for family, subset in split_sample_groups(chosen).items():
                kind, _ = SAMPLE_FAMILIES[family]
                unit = configured_units.get(family, '').strip()
                if not draft and not unit:
                    raise ValueError(f'{family}: cần xác nhận đơn vị kết quả.')
                if not draft and family in ('NS', 'NM', 'NT'):
                    raise ValueError(f'{family}: mẫu Word khí hiện tại không phù hợp SOP nền nước.')
                # Kết quả phiên có thể chưa lưu Ngày PT theo từng dòng. Khi người
                # dùng đã xác nhận ngày ở phần đầu biên bản, điền ngày ấy cho
                # các ô còn trống, rồi đồng bộ TG giải hấp = Ngày PT.
                subset = subset.copy()
                if analysis_date.strip():
                    subset.loc[subset['Ngày PT'].map(_as_text).eq(''), 'Ngày PT'] = analysis_date.strip()
                grouped_rows = number_tubes_and_vials(subset, insert_blank=True)
                group_hints = suggest_metadata(subset, 'excel' if is_excel else 'session')
                header = {'loai_mau': kind, 'ma_me': batch, 'nguoi_phan_tich': analyst,
                          'ngay_phan_tich': analysis_date, 'phuong_phap': method,
                          'chi_tieu_phan_tich': group_hints['chi_tieu_phan_tich']}
                result = generate_bien_ban(template_bytes, header, chem_rows, grouped_rows, unit, draft=draft)
                try:
                    xlsx = generate_bien_ban_excel(header, grouped_rows, unit, draft=draft)
                except (ImportError, ModuleNotFoundError):
                    # Web cũ đã có bộ xuất Excel BB-KT; vẫn dùng được trên Streamlit
                    # nếu môi trường máy chủ chưa có artifact_tool.
                    if excel_exporter is None:
                        raise ValueError('Chưa có thư viện xuất Excel và thiếu bộ xuất dự phòng từ app.py.')
                    xlsx = excel_exporter(result)
                clean_batch = re.sub(r'[^A-Za-z0-9_.-]+', '_', batch.strip())[:55] or 'CHUA_CO_MA_ME'
                base_name = f'{"NHAP_" if draft else ""}Bien_Ban_{family}_{clean_batch}'
                created.append((family, base_name+'.docx', result, base_name+'.xlsx', xlsx,
                                len(grouped_rows), subset['Tên mẫu'].nunique()))
            archive = BytesIO()
            with ZipFile(archive, 'w', ZIP_DEFLATED) as zf:
                for _, word_name, word_data, excel_name, excel_data, _, _ in created:
                    zf.writestr(word_name, word_data)
                    zf.writestr(excel_name, excel_data)
            st.session_state['bbxl_files'] = (source_digest, form_digest, created, archive.getvalue())
            st.success(f'Đã tạo {len(created)} nền, mỗi nền có 1 Word + 1 Excel. Rà soát hồ sơ trước khi phát hành.')
        except Exception as exc:
            st.error('Chưa xuất được Word: ' + str(exc))
    saved = st.session_state.get('bbxl_files')
    if saved and saved[0] == source_digest and saved[1] == form_digest:
        _, _, files, archive = saved
        for family, word_name, word_data, excel_name, excel_data, count, sample_count in files:
            left, right = st.columns(2)
            with left:
                st.download_button(f'Tải Word {family} — {sample_count} mẫu + BLANK',
                                   data=word_data, file_name=word_name,
                                   mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                                   key='bbxl_download_word_'+family+'_'+source_digest)
            with right:
                st.download_button(f'Tải Excel {family} — {count} dòng',
                                   data=excel_data, file_name=excel_name,
                                   mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                                   key='bbxl_download_excel_'+family+'_'+source_digest)
        st.download_button('Tải tất cả Word + Excel (.zip)', data=archive,
                           file_name=f'Bien_Ban_Tach_Nen_{re.sub(r"[^A-Za-z0-9_.-]+", "_", batch)}.zip',
                           mime='application/zip', key='bbxl_download_zip_'+source_digest)
