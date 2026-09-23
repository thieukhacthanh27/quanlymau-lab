import streamlit as st
import pandas as pd
from datetime import datetime, timedelta, date
import re
import difflib
import io
import json
import unicodedata
import os
import html
import math
import hashlib
from pathlib import Path
from streamlit_gsheets import GSheetsConnection

# ==========================================
# 1. CẤU HÌNH GIAO DIỆN & THIẾT KẾ (UI/UX)
# ==========================================
st.set_page_config(page_title="GC HATICO - Lab GC", page_icon="🔬", layout="wide")

st.markdown("""
<style>
.stApp {background:#f3f6f8;color:#142c3b}
.stMainBlockContainer {padding-top:2rem;max-width:1500px;padding-bottom:4rem}
[data-testid="stSidebar"] {background:#102d38;border-right:0}
[data-testid="stSidebar"] h2,[data-testid="stSidebar"] p,[data-testid="stSidebar"] label {color:#d9e8e9!important}
[data-testid="stSidebar"] [data-testid="stRadio"] label {padding:8px 10px;border-radius:8px;margin-bottom:3px}
[data-testid="stSidebar"] [data-testid="stRadio"] label:has(input:checked) {background:#21515d}
[data-testid="stSidebar"] .stButton button,[data-testid="stSidebar"] [data-testid="stPopover"] button {background:#1c424f;color:#e4f2f1;border-color:#355865}
[data-testid="stSidebar"] .stButton button p {color:#e4f2f1!important}
[data-testid="stSidebar"] hr {border-color:#34515c}
.main-title {font-size:2rem!important;font-weight:750;color:#142c3b;margin-bottom:.3rem}
.sub-title {font-size:1.15rem;font-weight:700;margin:1rem 0}
.eyebrow {font-size:11px;letter-spacing:2px;font-weight:750;color:#63848b;text-transform:uppercase}
.hero {display:flex;align-items:center;justify-content:space-between;gap:24px;margin:8px 0 25px}
.hero h1 {font-size:34px;font-weight:750;margin:5px 0 8px;padding:0;letter-spacing:-1px}
.hero p {color:#6a7e89;margin:0;font-size:14px}
.date-pill {background:white;border:1px solid #dce6e9;border-radius:25px;padding:10px 17px;font-size:12px;white-space:nowrap;color:#537079}
.brand {font-size:24px;letter-spacing:1px;font-weight:800;color:white;margin-top:4px}
.brand em {color:#5bd2ba;font-style:normal}
.brand-sub {font-size:11px;letter-spacing:2px;color:#8aaab5;margin:4px 0 28px}
.kpi {background:white;border:1px solid #e0e8ec;border-radius:13px;padding:20px 23px;min-height:150px;border-top:3px solid var(--accent)}
.kpi-label {color:#607986;font-size:13px;font-weight:600}
.kpi-number {font-size:36px;font-weight:750;line-height:1.5;color:#173847}
.kpi-foot {font-size:11px;color:#82959d}
.section-title {font-size:17px;font-weight:700;color:#173847;margin:0 0 4px}
.section-note {font-size:12px;color:#7b9199;margin-bottom:17px}
.flow-row {display:flex;align-items:center;gap:14px;margin:12px 0;font-size:12px}
.flow-label {width:125px;color:#5e7783;flex-shrink:0}
.flow-track {height:7px;flex:1;background:#edf2f4;border-radius:8px;overflow:hidden}
.flow-fill {height:100%;background:#23a68f;border-radius:8px}
.flow-count {width:24px;font-weight:700;text-align:right;color:#173847}
.notice {padding:12px 14px;background:#fff8eb;border:1px solid #f4e6c9;border-radius:9px;color:#89662e;font-size:12px;margin:10px 0}
.stButton button,.stDownloadButton button {border-radius:8px;font-weight:600;border:1px solid #d7e4e8;min-height:41px}
.stButton button[kind="primary"],.stDownloadButton button[kind="primary"] {background:#168d7d;border-color:#168d7d;color:white}
[data-testid="stVerticalBlockBorderWrapper"]>div {border-radius:13px!important}
[data-testid="stDataFrame"],[data-testid="stDataEditor"] {border:1px solid #e1e9ed;border-radius:9px;overflow:hidden}
[data-testid="stExpander"] {background:white;border-radius:10px}
[data-testid="stMetric"] {background:white;border:1px solid #e1e9ed;padding:16px;border-radius:12px}
.stTabs [data-baseweb="tab-list"] {gap:24px;background:transparent}
.stTabs [aria-selected="true"] {color:#168d7d!important}
.stTabs [data-baseweb="tab-highlight"] {background:#168d7d}
.warning-box {background:#fff4e9;border-left:4px solid #d99944;padding:15px;border-radius:7px}
@media(max-width:760px){.hero{display:block}.date-pill{display:inline-block;margin-top:15px}.hero h1{font-size:27px}.kpi{min-height:120px;padding:15px}.stMainBlockContainer{padding:1rem}}
</style>
""", unsafe_allow_html=True)

SHEET_URL = "https://docs.google.com/spreadsheets/d/1F2wFnxboWTFWDMGUuBDRGB901a5EKgvazHxkCgBjjRU/edit?usp=sharing"

STATUSES = [
    "🔴 1. Chờ xử lý", "🟠 2. Đang xử lý mẫu", "🟡 3. Chờ chạy máy",
    "🔵 4. Đang chạy máy", "🟣 5. Đang tính số liệu", "🟢 6. Lưu kho", "⚫ 7. Đã tiêu hủy"
]

CHEM_SYSTEMS = ["GC-MS", "GC-FID", "Thermo", "Dùng chung"]
CHEM_TYPES = ["Chất chuẩn (IS/Surrogate)", "Dung môi", "Vật tư tiêu hao", "Khí chuẩn", "Khác"]
CHEM_STATUS = ["🟢 Còn nhiều", "🟡 Sắp hết", "🔴 Đã hết"]

# KHAI BÁO BIẾN CHO KHO HÓA CHẤT
U_GROUPS = {
    'Nồng độ dung dịch': {'g/L': 1000., 'mg/L': 1., 'µg/L': .001, 'ng/L': .000001, 'mg/mL': 1000., 'µg/mL': 1., 'ng/mL': .001},
    'Thể tích': {'L': 1., 'mL': .001, 'µL': .000001, 'm³': 1000.},
    'Khối lượng': {'g': 1., 'mg': .001, 'µg': .000001, 'ng': .000000001, 'kg': 1000.},
    'Nồng độ khí thực': {'g/m³': 1000., 'mg/m³': 1., 'µg/m³': .001, 'ng/m³': .000001},
    'Nồng độ khí chuẩn': {'mg/Nm³': 1., 'µg/Nm³': .001, 'ng/Nm³': .000001},
    'Hàm lượng khối lượng': {'mg/kg': 1., 'µg/kg': .001, 'ng/kg': .000001, 'ppm (m/m)': 1., 'ppb (m/m)': .001, '% (m/m)': 10000.}
}
CHEM_BASE = ['Hệ Máy', 'Phân Loại', 'Tên Hóa Chất', 'Số Lô (Lot)', 'Ngày Mở Nắp', 'Hạn Sử Dụng', 'Tình Trạng Kho', 'Ghi Chú']
CHEM_EXTRA = ['ID Nguồn', 'STT Nguồn', 'CAS', 'Nhà Sản Xuất', 'Nồng Độ', 'Đơn Vị Nồng Độ', 'Độ Tinh Khiết (%)', 'Quy Cách Gốc', 'Lượng Quy Cách', 'Đơn Vị Quy Cách', 'Tình Trạng Gốc', 'HSD Gốc', 'Bảo Quản', 'Nguồn PDF', 'Trang PDF', 'Dòng PDF', 'Cần Kiểm Tra']

# ==========================================
# 2. HÀM TIỆN ÍCH KHO HÓA CHẤT NÂNG CAO
# ==========================================
def u_text(v):
    return '' if v is None or pd.isna(v) else str(v).strip()

def u_canonical(unit):
    text = u_text(unit).replace('μ', 'µ').replace('ug', 'µg').replace('ul', 'µL').replace('m3', 'm³').replace('ml', 'mL').replace(' / ', '/')
    return {'ng/ml': 'ng/mL', 'ug/L': 'µg/L', 'ug/ml': 'µg/mL'}.get(text, text)

def u_convert(value, source, target):
    value = float(value)
    if not math.isfinite(value): raise ValueError('Giá trị phải hữu hạn.')
    source, target = u_canonical(source), u_canonical(target)
    for group in U_GROUPS.values():
        if source in group and target in group: return value * group[source] / group[target]
    raise ValueError(f'Không tự quy đổi {source} → {target}: khác đại lượng hoặc chưa xác định cơ sở đơn vị.')

def u_expiry(text):
    text = u_text(text)
    m = re.fullmatch(r'(\d{1,2})/(\d{1,2})/(\d{4})', text)
    if not m: return None, ('Thiếu/ngày không đầy đủ' if text else 'Chưa có HSD')
    a, b, y = map(int, m.groups())
    if a <= 12 and b <= 12 and a != b: return None, 'Ngày mơ hồ D/M hay M/D: cần xác nhận'
    try: return date(y, b, a) if a > 12 or a == b else date(y, a, b), ''
    except ValueError: return None, 'Ngày không hợp lệ'

def u_cas_valid(cas):
    digits = cas.replace('-', '')
    return bool(re.fullmatch(r'\d{2,7}-\d{2}-\d', cas)) and sum(int(v) * (i + 1) for i, v in enumerate(reversed(digits[:-1]))) % 10 == int(digits[-1])

def read_stock_pdf(data, filename):
    import pdfplumber
    digest = hashlib.sha256(data).hexdigest()
    records = []; system = ''; kind = 'Chất chuẩn phân tích'
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page_no, page in enumerate(pdf.pages, 1):
            for table in page.extract_tables():
                for row_no, cells in enumerate(table, 1):
                    cells = [u_text(x).replace('\n', ' ') for x in cells]
                    if len(cells) != 9: continue
                    stt, name, formula, maker, pack, state, expiry, storage, unlabelled = cells
                    joined = ' '.join(cells).upper()
                    if 'CHO MÁY GCMS THERMO' in joined: system = 'Thermo'; kind = 'Chất chuẩn phân tích'; continue
                    if 'GCMS AGILENT' in joined: system = 'Dùng chung'; kind = 'Chất chuẩn phân tích'; continue
                    if 'ĐỒNG HÀNH VÀ NỘI CHUẨN' in joined: kind = 'Chất chuẩn (IS/Surrogate)'; continue
                    if name.strip().upper() in ['OPPS', 'OCPS', 'PCBS', 'PHENOL']: kind = 'Chất chuẩn phân tích'; continue
                    if not name or name in ['Tên hóa chất', 'Tên Hóa Chất'] or not (maker or pack or state): continue
                    text = name + ' ' + formula
                    warnings = []
                    cas_values = re.findall(r'\b\d{2,7}\s*-\s*\d{2}\s*-\s*\d\b', text)
                    cas_values = list(dict.fromkeys(re.sub(r'\s', '', v) for v in cas_values))
                    if any(not u_cas_valid(c) for c in cas_values): warnings.append('CAS không khớp số kiểm tra; đối chiếu COA')
                    values = re.findall(r'(\d+(?:[.,]\d+)?)\s*(mg|ug|µg|μg|ng)\s*/\s*(ml|mL|L|l)\b', text)
                    concentrations = {(float(n.replace(',', '.')), u_canonical(a + '/' + ('mL' if b.lower() == 'ml' else 'L'))) for n, a, b in values}
                    conc, unit = next(iter(concentrations)) if len(concentrations) == 1 else (None, '')
                    if len(concentrations) > 1: warnings.append('Nhiều nồng độ trong một dòng; không tự gộp')
                    purity = re.search(r'(\d+(?:[.,]\d+)?)\s*%', text)
                    pure = float(purity.group(1).replace(',', '.')) if purity else None
                    pm = re.search(r'(\d+(?:[.,]\d+)?)\s*(mg|ug|µg|g|ml|mL|L)\b', pack)
                    amount = float(pm.group(1).replace(',', '.')) if pm else None
                    pack_unit = u_canonical(pm.group(2)) if pm else ''
                    exp, date_note = u_expiry(expiry)
                    if date_note: warnings.append(date_note)
                    if len(name) < 3: warnings.append('Tên hóa chất chưa đầy đủ')
                    if re.search(r'components|each', text, re.I): warnings.append('Hỗn hợp: giữ nguyên thành phần; nồng độ có thể là mỗi chất')
                    status = '🔴 Đã hết' if 'hết' in state.lower() else 'Chưa xác nhận lượng tồn'
                    notes = ('Công thức/thành phần: ' + formula if formula else '')
                    notes += '; Cột cuối không có tiêu đề: ' + unlabelled if unlabelled else ''
                    records.append(dict(zip(CHEM_BASE, [system or 'Dùng chung', kind, name, '', None, exp, status, notes])) | {
                        'ID Nguồn': digest + f':{page_no}:{row_no}', 'STT Nguồn': stt, 'CAS': '; '.join(cas_values), 'Nhà Sản Xuất': maker,
                        'Nồng Độ': conc, 'Đơn Vị Nồng Độ': unit, 'Độ Tinh Khiết (%)': pure, 'Quy Cách Gốc': pack,
                        'Lượng Quy Cách': amount, 'Đơn Vị Quy Cách': pack_unit, 'Tình Trạng Gốc': state, 'HSD Gốc': expiry,
                        'Bảo Quản': storage, 'Nguồn PDF': filename, 'Trang PDF': page_no, 'Dòng PDF': row_no, 'Cần Kiểm Tra': '; '.join(warnings)})
    if not records: raise ValueError('Không tìm thấy bảng 9 cột theo PDF kho GC đã cung cấp. Không hỗ trợ PDF scan/cấu trúc khác.')
    df = pd.DataFrame(records)
    duplicate = df.duplicated(['Tên Hóa Chất', 'Nhà Sản Xuất', 'Quy Cách Gốc', 'HSD Gốc'], keep=False)
    df.loc[duplicate, 'Cần Kiểm Tra'] = df.loc[duplicate, 'Cần Kiểm Tra'].map(lambda x: x + '; Có dòng có thể trùng trong PDF: đối chiếu từng lọ, không tự cộng kho')
    return df

def stock_frame(df):
    if not df.empty and 'Tên Hóa Chất' not in df: raise ValueError('QuanLyHoaChat thiếu cột Tên Hóa Chất. Không ghi đè cấu trúc.')
    df = df.copy()
    for c in CHEM_BASE + CHEM_EXTRA:
        if c not in df: df[c] = None if c in ['Ngày Mở Nắp', 'Hạn Sử Dụng', 'Nồng Độ', 'Độ Tinh Khiết (%)', 'Lượng Quy Cách'] else ''
    for c in ['Ngày Mở Nắp', 'Hạn Sử Dụng']: df[c] = pd.to_datetime(df[c], errors='coerce').dt.date
    for c in ['Nồng Độ', 'Độ Tinh Khiết (%)', 'Lượng Quy Cách']: df[c] = pd.to_numeric(df[c], errors='coerce')
    return df

def stock_serial(df):
    df = stock_frame(df)
    for c in ['Ngày Mở Nắp', 'Hạn Sử Dụng']: df[c] = pd.to_datetime(df[c], errors='coerce').dt.strftime('%Y-%m-%d')
    return df.fillna('')

def stock_fingerprint(df):
    return hashlib.sha256(stock_serial(df).astype(str).to_json(orient='split', force_ascii=False).encode()).hexdigest()

def stock_merge(base, incoming):
    existing = set(base.get('ID Nguồn', pd.Series(dtype=str)).dropna().astype(str)) - {''}
    ids = incoming['ID Nguồn'].astype(str)
    if ids.duplicated().any() or any(x in existing for x in ids if x): raise ValueError('Đã nhập dòng nguồn PDF/CSV này. Chưa lưu thêm dòng nào.')
    return pd.concat([base, incoming], ignore_index=True)

# ==========================================
# 3. KẾT NỐI DATABASE CHÍNH
# ==========================================
DEMO_MODE = os.environ.get("LAB_DEMO", "0") == "1"
class DemoConnection:
    def read(self, spreadsheet=None, worksheet=None, **kwargs):
        key = "demo_store_" + (worksheet or "samples")
        if key not in st.session_state:
            if worksheet == "QuanLyHoaChat":
                rows = []
                for name, offset, system in [("VOCs mix • DEMO", 18, "GC-MS"), ("Toluene-d8 • DEMO", 160, "GC-MS"), ("PCB mix • DEMO", -8, "Thermo"), ("Methanol • DEMO", 300, "Dùng chung")]:
                    rows.append({"Tên Hóa Chất":name,"Hệ Máy":system,"Phân Loại":"Chất chuẩn (IS/Surrogate)","Số Lô (Lot)":"DEMO-2026","Ngày Mở Nắp":"","Hạn Sử Dụng":(datetime.now()+timedelta(days=offset)).strftime("%Y-%m-%d"),"Tình Trạng Kho":"🟢 Còn nhiều","Ghi Chú":"Dữ liệu minh họa","Nồng Độ":100.0,"Đơn Vị Nồng Độ":"µg/mL"})
                data = pd.DataFrame(rows)
            elif worksheet:
                data = pd.DataFrame(columns=["Nền Mẫu", "Tên Chất", "MDL", "LOQ", "Đơn Vị"])
            else:
                rows = []
                for i in range(36):
                    rows.append({"Mã Mẫu":f"{'NS' if i%3 else 'KT'}-DEMO-{i+1:03}","Tên Mẻ":f"DEMO.2026.{i//12+1:03}","Nền Mẫu":"Nước" if i%3 else "Khí","Chỉ Tiêu":"VOCs; Benzen; Toluen" if i%2 else "OCP; PCB","Trạng Thái":STATUSES[i%7],"Người Giữ":["Thành","KTV 02","KTV 03"][i%3],"Ghi Chú":"Mẫu minh họa","Giờ Nhận":datetime.now()-timedelta(days=i%4,hours=i%5)})
                data = pd.DataFrame(rows)
            st.session_state[key] = data
        return st.session_state[key].copy(deep=True)
    def update(self, spreadsheet=None, worksheet=None, data=None, **kwargs):
        st.session_state["demo_store_"+(worksheet or "samples")] = data.copy(deep=True)
        return data

conn = DemoConnection() if DEMO_MODE else st.connection("gsheets", type=GSheetsConnection)

def load_data():
    df = conn.read(spreadsheet=SHEET_URL, ttl=0)
    if df.empty or len(df.columns) == 0 or "Mã Mẫu" not in df.columns:
        df = pd.DataFrame(columns=["Mã Mẫu", "Tên Mẻ", "Nền Mẫu", "Chỉ Tiêu", "Trạng Thái", "Người Giữ", "Ghi Chú", "Giờ Nhận"])
        conn.update(spreadsheet=SHEET_URL, data=df)
    df['Giờ Nhận'] = pd.to_datetime(df['Giờ Nhận'], errors='coerce')
    return df

def save_data(df):
    df_save = df.copy()
    df_save['Giờ Nhận'] = df_save['Giờ Nhận'].dt.strftime('%Y-%m-%d %H:%M:%S')
    conn.update(spreadsheet=SHEET_URL, data=df_save)
    st.cache_data.clear()

def load_limit_config():
    try:
        df_limit = conn.read(spreadsheet=SHEET_URL, worksheet="CauHinh_MDL_LOQ", ttl=0) 
        return df_limit
    except:
        return pd.DataFrame(columns=["Nền Mẫu", "Tên Chất", "MDL", "LOQ", "Đơn Vị"])

def load_chemical_data():
    try:
        df = stock_frame(conn.read(spreadsheet=SHEET_URL, worksheet='QuanLyHoaChat', ttl=0))
        st.session_state.chem_base = stock_fingerprint(df)
        st.session_state.pop('chem_error', None)
        return df
    except Exception as exc:
        st.session_state.chem_error = f'Không đọc được kho: {exc}. Kiểm tra worksheet QuanLyHoaChat rồi tải lại.'
        return stock_frame(pd.DataFrame())

def save_chemical_data(df_chem):
    if st.session_state.get('chem_error'): raise ValueError(st.session_state.chem_error)
    current = stock_frame(conn.read(spreadsheet=SHEET_URL, worksheet='QuanLyHoaChat', ttl=0))
    if stock_fingerprint(current) != st.session_state.get('chem_base'): raise ValueError('Kho đã thay đổi trên Sheets; tải lại trước khi lưu để tránh ghi đè.')
    
    for column in ['Nồng Độ', 'Độ Tinh Khiết (%)', 'Lượng Quy Cách']:
        values = pd.to_numeric(df_chem[column], errors='coerce')
        if (values.dropna() < 0).any(): raise ValueError(f'{column} không được âm.')
    if df_chem['Tên Hóa Chất'].fillna('').str.strip().eq('').any(): raise ValueError('Tên hóa chất không được trống.')
    
    saved = stock_serial(df_chem)
    conn.update(spreadsheet=SHEET_URL, worksheet='QuanLyHoaChat', data=saved)
    st.session_state.chem_base = stock_fingerprint(saved)
    st.cache_data.clear()

if "df" not in st.session_state:
    st.session_state.df = load_data()
if "df_limit" not in st.session_state:
    st.session_state.df_limit = load_limit_config()
if "df_chem" not in st.session_state:
    st.session_state.df_chem = load_chemical_data()

# ==========================================
# 4. CÁC HÀM BỔ SUNG & XỬ LÝ SỐ LIỆU
# ==========================================
def clean_str(value):
    return "" if pd.isna(value) else str(value).strip()

def parse_sample_matrix(sample_name):
    name_upper = str(sample_name).upper()
    if 'KT' in name_upper: return 24.0, 'KT', 'Khí'
    elif 'KXQ' in name_upper: return 4.0, 'KXQ', 'Khí'
    elif 'KLV' in name_upper: return 4.0, 'KLV', 'Khí'
    elif any(k in name_upper for k in ['NS', 'NT', 'NM', 'NN']): return 1.0, 'NS', 'Nước' 
    return None, None, None

def get_dynamic_surrogate_expected(c_do):
    levels = [1.0, 2.0, 4.0, 5.0, 6.0, 8.0, 10.0, 20.0, 25.0, 50.0, 100.0]
    valid_levels = []
    for lvl in levels:
        rec = (c_do / lvl) * 100.0
        if 70 <= rec <= 130:
            valid_levels.append((lvl, abs(100 - rec)))
    if valid_levels:
        valid_levels.sort(key=lambda x: x[1])
        return valid_levels[0][0]
    return min(levels, key=lambda x: abs(x - c_do))

def get_limit_info(compound_name, nen_mau):
    df_limit = st.session_state.df_limit
    if not df_limit.empty and compound_name and nen_mau:
        search_nen = "NS" if nen_mau in ['NT', 'NM', 'NN'] else nen_mau
        
        mask_exact = (df_limit["Nền Mẫu"].astype(str).str.upper() == search_nen.upper()) & \
                     (df_limit["Tên Chất"].astype(str).str.lower() == str(compound_name).lower())
        match = df_limit[mask_exact]
        
        if match.empty:
            all_comps = df_limit[df_limit["Nền Mẫu"].astype(str).str.upper() == search_nen.upper()]["Tên Chất"].astype(str).tolist()
            close_matches = difflib.get_close_matches(str(compound_name).lower(), [c.lower() for c in all_comps], n=1, cutoff=0.7)
            if close_matches:
                match = df_limit[(df_limit["Nền Mẫu"].astype(str).str.upper() == search_nen.upper()) & (df_limit["Tên Chất"].astype(str).str.lower() == close_matches[0])]

        if not match.empty:
            def parse_val(col_name):
                if col_name in match.columns:
                    val = str(match[col_name].values[0]).replace(',', '.')
                    try: return float(val)
                    except: return None
                return None

            mdl_val = parse_val("MDL")
            loq_val = parse_val("LOQ")
            unit = str(match["Đơn Vị"].values[0]) if "Đơn Vị" in match.columns else ""
            if pd.isna(unit) or unit == 'nan': unit = ""
            
            return mdl_val, loq_val, unit
    return None, None, ""

def convert_unit_value(value, from_unit, to_unit):
    """Hàm tự động quy đổi đơn vị đo lường cơ bản (Giữ lại cho module cũ)"""
    if pd.isna(value) or to_unit == "Mặc định" or not from_unit: return value
    try: val = float(value)
    except: return value

    f_u = str(from_unit).strip().lower()
    t_u = str(to_unit).strip().lower()

    if f_u == t_u: return val

    conversion_factors = {
        ('mg/l', 'µg/l'): 1000.0, ('mg/l', 'ppb'): 1000.0,
        ('µg/l', 'mg/l'): 0.001, ('ppb', 'mg/l'): 0.001,
        ('ppm', 'ppb'): 1000.0, ('ppb', 'ppm'): 0.001,
        ('mg/m3', 'µg/m3'): 1000.0, ('µg/m3', 'mg/m3'): 0.001,
        ('mg/nm3', 'µg/nm3'): 1000.0, ('µg/nm3', 'mg/nm3'): 0.001,
    }

    factor = conversion_factors.get((f_u, t_u))
    if factor: return val * factor
    return val

def u_controls(key):
    with st.expander('Đơn vị đầu vào và đầu ra',expanded=True):
        liquid=list(U_GROUPS['Nồng độ dung dịch'])
        src=st.selectbox('Đơn vị C đo và Csurr đo',liquid,index=liquid.index('µg/mL'),key=key+'src')
        ref=st.selectbox('Đơn vị Csurr trước / các mức C tự gợi ý',liquid,index=liquid.index('µg/mL'),key=key+'ref')
        water=st.selectbox('Đơn vị kết quả nước',liquid,index=liquid.index('µg/L'),key=key+'water')
        gas_mode=st.selectbox('Cách tính mẫu khí',['Công thức gốc: đơn vị kết quả do SOP xác định','Từ nồng độ dịch chiết và thể tích'],key=key+'mode')
        basis=st.selectbox('Cơ sở thể tích khí',['Thể tích thực','Thể tích đã quy về điều kiện chuẩn theo SOP'],key=key+'basis')
        gas_units=list(U_GROUPS['Nồng độ khí chuẩn' if basis.startswith('Thể tích đã') else 'Nồng độ khí thực'])
        gas_base=st.selectbox('Đơn vị kết quả khí trước đổi (theo SOP)',gas_units,index=gas_units.index('mg/Nm³' if basis.startswith('Thể tích đã') else 'mg/m³'),key=key+'base'+basis)
        gas_out=st.selectbox('Đơn vị xuất kết quả khí',gas_units,index=gas_units.index('µg/Nm³' if basis.startswith('Thể tích đã') else 'µg/m³'),key=key+'out'+basis)
        extraction=st.number_input('Thể tích dịch chiết/giải hấp (mL)',min_value=.000001,value=1.,key=key+'extract')
        st.caption('Nếu chọn khí chuẩn, V khí phải là thể tích đã được quy chuẩn từ trước theo SOP; app không tự giả định nhiệt độ/áp suất.')
    confirmed=st.checkbox('Đã xác nhận đơn vị nguồn, đơn vị kết quả và thông số theo SOP',key=key+'confirmed')
    cfg=dict(basis=basis,confirmed=confirmed,src=src,ref=ref,water=water,gas_mode=gas_mode,gas_base=gas_base,gas_out=gas_out,extraction=extraction)
    previous=st.session_state.get('unit_cfg')
    if previous is not None and previous!=cfg:st.session_state.results_stale=True
    st.session_state.unit_cfg=cfg
    return cfg

def u_result(raw,volume,mdl,loq,limit_unit,kind,recovery,cfg):
    if not math.isfinite(float(raw)) or raw<0:raise ValueError('C đo phải là số không âm hợp lệ.')
    if not math.isfinite(float(recovery)) or recovery<=0:raise ValueError('Recovery phải lớn hơn 0.')
    if kind=='Nước':
        out=cfg['water'];value=u_convert(raw,cfg['src'],out)*100/recovery
    else:
        if volume<=0:raise ValueError('Thể tích khí phải >0 L.')
        out=cfg['gas_out']
        if cfg['gas_mode'].startswith('Từ'):
            mass_mg=u_convert(raw,cfg['src'],'mg/L')*cfg['extraction']/1000
            value=u_convert(mass_mg/(volume/1000)*100/recovery,'mg/Nm³' if cfg.get('basis','').startswith('Thể tích đã') else 'mg/m³',out)
        else:
            value=u_convert(u_convert(raw,cfg['src'],cfg['ref'])/volume*100/recovery,cfg['gas_base'],out)
    threshold=mdl if kind=='Khí' else loq
    tag='MDL' if kind=='Khí' else 'LOQ'
    if threshold is not None and pd.notna(threshold):
        threshold=u_convert(float(threshold),limit_unit,out)
    report=f'{value:.8g}'
    if raw==0:report='KPH (quy ước bản gốc)'
    elif threshold is not None and pd.notna(threshold) and value<threshold:report=f'KPH (< MDL {threshold:g} {out})' if kind=='Khí' else f'< LOQ ({threshold:g} {out})'
    return report,value,out

def evaluate_result(raw_conc,v_param,mdl_val,loq_val,unit,loai_mau,recovery=100.0):
    cfg=st.session_state.get('unit_cfg')
    if not cfg or not cfg.get('confirmed'):
        # Fallback to old behavior if unit config is not active/confirmed
        if pd.isna(raw_conc) or raw_conc <= 0: return "KPH"
        if loai_mau == 'Khí':
            c_thuc_val = (raw_conc * 1.0) / v_param * (100.0 / recovery)
            if mdl_val is not None and c_thuc_val < mdl_val: return f"KPH (< MDL: {mdl_val} {unit})"
        elif loai_mau == 'Nước':
            c_thuc_val = raw_conc * (100.0 / recovery)
            if loq_val is not None and c_thuc_val < loq_val: return f"< LOQ ({loq_val} {unit})"
        else:
            c_thuc_val = raw_conc
        return f"{round(c_thuc_val, 4)}"
    return u_result(raw_conc,v_param,mdl_val,loq_val,unit,loai_mau,recovery,cfg)[0]

def unit_utility():
    st.markdown("<h1 class='main-title'>🔄 Quy đổi đơn vị</h1>", unsafe_allow_html=True)
    with st.container(border=True):
        group=st.selectbox('Đại lượng',list(U_GROUPS))
        units=list(U_GROUPS[group])
        a=st.selectbox('Từ đơn vị',units);b=st.selectbox('Sang đơn vị',units,index=min(1,len(units)-1))
        v=st.number_input('Giá trị',value=1.0)
        st.metric('Giá trị sau đổi',f'{u_convert(v,a,b):.10g} {b}')
        st.caption('1 µg/mL = 1 mg/L = 1000 µg/L. Độ tinh khiết (%) không phải nồng độ dung dịch. Khối lượng và thể tích cần khối lượng riêng để đổi.')

def lab_norm(value):
    if value is None or pd.isna(value): return ''
    return ''.join(c for c in unicodedata.normalize('NFD',str(value).lower().replace('đ','d')) if unicodedata.category(c)!='Mn').strip()

def lab_local_answer(question, df):
    q = lab_norm(question)
    
    if any(x in q for x in ['c ban dau', 'thu hoi', 'r%', 'csurr', 'duong chuan']):
        return "💡 **R(%) = (C_surr_sau / C_surr_trước) × 100**.\nHệ thống tự động chọn C ban đầu từ danh sách mức cố định sao cho R(%) rơi vào 70–130%."
    if any(x in q for x in ['cong thuc', 'cach tinh']):
        return "💡 **Công thức SOP:**\n- Mẫu Nước: C = C đo × 100/R.\n- Mẫu Khí: C = C đo ÷ V × 100/R.\nHệ thống tự động gợi ý V=24L cho mẫu Khí (KT) và V=4L cho mẫu Xung quanh (KXQ/KLV)."
    if any(x in q for x in ['bien ban', 'huong dan', 'cach nhap', 'loi', 'mdl', 'loq']):
        return "💡 **Hướng dẫn:**\n- Tiếp nhận mẫu ở Tab 2.\n- Tính kết quả tự động ở GC-MS.\n- Tính tay & Pha chuẩn ở Công cụ Phân tích.\n- Lập Biên bản & In tem ở Báo cáo & Lập Biên bản."

    filters = []
    code = re.search(r'\b(?:ns|nt|nm|nn|kt|kxq|kkxq|klv)[.\-]?\d[\w.\-]*', q)
    if code: filters.append(('Mã Mẫu', code.group()))
    batch = re.search(r'\b\d{4}\.\d{2}\.\d{3}\b', q)
    if batch: filters.append(('Tên Mẻ', batch.group()))
    status_words = [('cho chay', '3'), ('dang chay', '4'), ('cho xu ly', '1'), ('dang xu ly', '2'), ('tinh so lieu', '5'), ('luu kho', '6'), ('tieu huy', '7')]
    for phrase, status in status_words:
        if phrase in q: filters.append(('Trạng Thái', status))
    if 'mau nuoc' in q: filters.append(('Nền Mẫu', 'Nước'))
    elif 'mau khi' in q: filters.append(('Nền Mẫu', 'Khí'))

    res_df = df.copy()
    for field, val in filters:
        if field == 'Trạng Thái':
            res_df = res_df[res_df[field].str.contains(val, na=False)]
        else:
            res_df = res_df[res_df[field].astype(str).str.lower().str.contains(val.lower(), na=False)]

    if not filters and not any(x in q for x in ['tong', 'bao nhieu', 'thong ke', 'danh sach', 'nhom theo', 'cua ai', 'ai dang giu']):
        return "🤔 Mình chưa hiểu ý bạn. Bạn có thể hỏi cụ thể:\n- *Tra mẫu NS-150826-004*\n- *Có bao nhiêu mẫu chờ chạy?*\n- *Ai đang giữ nhiều mẫu nhất?*"

    if any(phrase in q for phrase in ['theo nguoi', 'ai dang giu', 'cua ai']):
        users = res_df['Người Giữ'].replace("", "Chưa phân công").dropna().value_counts()
        ans = f"👥 **Phân bổ mẫu ({len(res_df)} mẫu):**\n"
        for user, count in users.items(): ans += f"- **{user}**: {count} mẫu\n"
        return ans
    elif 'theo trang thai' in q:
        stats = res_df['Trạng Thái'].value_counts()
        ans = f"📊 **Trạng thái ({len(res_df)} mẫu):**\n"
        for st_name, count in stats.items(): ans += f"- {st_name}: {count} mẫu\n"
        return ans

    if len(res_df) == 0:
        return "❌ Không tìm thấy mẫu nào khớp với dữ liệu bạn tìm kiếm."
    elif len(res_df) == 1:
        info = res_df.iloc[0]
        return f"🔍 **Đã tìm thấy mẫu {info['Mã Mẫu']}:**\n- **Thuộc mẻ:** {info['Tên Mẻ']}\n- **Nền:** {info['Nền Mẫu']}\n- **Trạng thái:** {info['Trạng Thái']}\n- **Người giữ:** {info['Người Giữ']}\n- **Chỉ tiêu:** {info['Chỉ Tiêu']}"
    else:
        ans = f"📋 **Tìm thấy {len(res_df)} mẫu khớp yêu cầu:**\n"
        for _, row in res_df.head(10).iterrows():
            ans += f"- **{row['Mã Mẫu']}** (Mẻ: {row['Tên Mẻ']} | Trạng thái: {row['Trạng Thái'][:4]})\n"
        if len(res_df) > 10: ans += "*... và nhiều mẫu khác.*"
        return ans

# ==========================================
# 5. THANH ĐIỀU HƯỚNG BÊN TRÁI (SIDEBAR)
# ==========================================
st.sidebar.markdown("<div class='brand'>HATICO<em> / LAB</em></div><div class='brand-sub'>GC · LABORATORY WORKSPACE</div>",unsafe_allow_html=True)
if DEMO_MODE:
    st.sidebar.caption("🧪 BẢN THỬ · Dữ liệu minh họa")
else:
    st.sidebar.caption("Nguồn dữ liệu: Google Sheets")

def go_page(page):
    st.session_state["nav_page"] = page

menu = st.sidebar.radio("📌 ĐIỀU HƯỚNG CHÍNH", [
    "🏠 Trang chủ (Tổng quan)", 
    "📥 Quản lý Tiếp nhận", 
    "⚙️ Vận hành GC-MS",
    "🔥 Vận hành GC-FID",
    "🧬 Vận hành Thermo",
    "🧮 Tiện ích Phân tích",
    "📝 Báo cáo & Lập Biên bản",
    "🧪 Kiểm soát Hóa chất",
    "🔄 Quy đổi đơn vị",
    "⚙️ Cấu hình Hệ thống"
], key="nav_page", label_visibility="collapsed")

st.sidebar.divider()

if st.sidebar.button("🔄 Cập nhật hệ thống", use_container_width=True):
    st.cache_data.clear()
    if 'results' in st.session_state:
        st.session_state.results_stale = True
    st.session_state.df = load_data()
    st.session_state.df_limit = load_limit_config() 
    st.session_state.df_chem = load_chemical_data()
    st.rerun()

st.sidebar.divider()

with st.sidebar.popover("💬 Trợ lý tra cứu", use_container_width=True):
    st.markdown("👋 **Xin chào! Mình là Trợ lý LIMS.**")
    st.caption("Tra cứu tiến độ mẫu, mẻ, phân công hoặc hỏi về SOP.")
    
    if "chat_history" not in st.session_state:
        st.session_state.chat_history = []
        
    chat_container = st.container(height=350)
    with chat_container:
        for msg in st.session_state.chat_history:
            if msg["role"] == "user":
                st.markdown(f"👤 **Bạn:** {msg['content']}")
            else:
                st.info(f"🤖 **Bot:**\n{msg['content']}")
                
    with st.form("chat_form", clear_on_submit=True):
        user_input = st.text_input("Gõ câu hỏi...", placeholder="VD: Mẫu NS-01 đang ở đâu?")
        submit_btn = st.form_submit_button("Gửi")
        if submit_btn and user_input.strip():
            st.session_state.chat_history.append({"role": "user", "content": user_input})
            ans = lab_local_answer(user_input, st.session_state.df)
            st.session_state.chat_history.append({"role": "bot", "content": ans})
            st.rerun()
            
    if st.session_state.chat_history:
        if st.button("🧹 Xóa hội thoại", use_container_width=True):
            st.session_state.chat_history = []
            st.rerun()

df_current = st.session_state.df.copy()
df_current["Ngày Nhận"] = df_current["Giờ Nhận"].dt.date
today_date = datetime.today().date()

# ==========================================
# 6. GIAO DIỆN CÁC TRANG
# ==========================================

if DEMO_MODE:
    st.caption("🧪 CHẾ ĐỘ DEMO — Mẫu và hóa chất minh họa; các thay đổi chỉ lưu trong phiên thử.")

if menu == "🏠 Trang chủ (Tổng quan)":
    st.markdown(f"""<div class="hero"><div><div class="eyebrow">Không gian vận hành / Tổng quan</div><h1>Một ngày làm việc hiệu quả.</h1><p>Theo dõi mẫu, kiểm soát tiến độ và chuẩn bị cho lượt chạy tiếp theo.</p></div><div class="date-pill">◷ &nbsp; {today_date.strftime('%d / %m / %Y')}</div></div>""",unsafe_allow_html=True)
    active = ~df_current["Trạng Thái"].isin(STATUSES[-2:])
    received = int((df_current["Ngày Nhận"] == today_date).sum())
    ready = int(df_current["Trạng Thái"].eq(STATUSES[2]).sum())
    pending = int(((df_current["Ngày Nhận"] < today_date) & active).sum())
    done = int((~active).sum())
    
    cards=[("TIẾP NHẬN HÔM NAY",received,"Theo ngày nhận mẫu", "#239f8d"),("SẴN SÀNG CHẠY MÁY",ready,"Đã chuyển sang chờ chạy", "#5d8fbe"),("MẪU CÒN TỒN",pending,"Nhận trước hôm nay, chưa lưu/hủy", "#dbab5b"),("ĐÃ LƯU / TIÊU HỦY",done,"Trên toàn bộ danh sách", "#8b82be")]
    for col,(label,n,foot,color) in zip(st.columns(4),cards):
        col.markdown(f'<div class="kpi" style="--accent:{color}"><div class="kpi-label">{label}</div><div class="kpi-number">{n:02}</div><div class="kpi-foot">{foot}</div></div>',unsafe_allow_html=True)
    st.write("")
    
    shortcuts=[("＋ Tiếp nhận mẫu","📥 Quản lý Tiếp nhận"),("▷ Xử lý GC-MS","⚙️ Vận hành GC-MS"),("▤ Lập biên bản","📝 Báo cáo & Lập Biên bản"),("⇄ Quy đổi đơn vị","🔄 Quy đổi đơn vị")]
    for col,(label,target) in zip(st.columns(4),shortcuts):
        col.button(label,on_click=go_page,args=(target,),use_container_width=True)
        
    left,right=st.columns([1.5,1],gap="large")
    with left,st.container(border=True):
        st.markdown('<div class="section-title">Tiến độ phòng lab</div><div class="section-note">Phân bố mẫu theo trạng thái hiện tại</div>',unsafe_allow_html=True)
        for status in STATUSES[:5]:
            count=int(df_current["Trạng Thái"].eq(status).sum())
            label=status.split(". ",1)[-1]
            st.markdown(f'<div class="flow-row"><span class="flow-label">{label}</span><div class="flow-track"><div class="flow-fill" style="width:{100*count/max(len(df_current),1):.1f}%"></div></div><span class="flow-count">{count}</span></div>',unsafe_allow_html=True)
            
    with right,st.container(border=True):
        st.markdown('<div class="section-title">Cần chú ý</div><div class="section-note">Kiểm tra trước khi phân tích</div>',unsafe_allow_html=True)
        chem=st.session_state.df_chem
        dates=pd.to_datetime(chem['Hạn Sử Dụng'],errors='coerce')
        expired=int((dates.dt.date<today_date).sum())
        soon=int(((dates.dt.date>=today_date)&(dates.dt.date<=today_date+timedelta(days=30))).sum())
        if st.session_state.get('chem_error'):
            st.warning("Chưa đọc được kho hóa chất. Mở kho để kiểm tra kết nối.")
        else:
            st.markdown(f'<div class="notice">◷ &nbsp; <b>{soon} hóa chất</b> hết hạn trong 30 ngày · <b>{expired}</b> đã hết hạn</div>',unsafe_allow_html=True)
            st.caption(f"{int(dates.isna().sum())} dòng chưa xác định hạn sử dụng.")
        unassigned=int((active & df_current['Người Giữ'].fillna('').str.strip().eq('')).sum())
        st.caption(f"{unassigned} mẫu đang xử lý chưa có người phụ trách.")
        st.button("Mở kho hóa chất →",on_click=go_page,args=("🧪 Kiểm soát Hóa chất",),use_container_width=True)
        
    st.markdown('<div class="section-title">Danh sách công việc</div>',unsafe_allow_html=True)
    st.caption("Tra cứu, phân công và cập nhật trạng thái mẫu tại một nơi.")
    
    with st.container(border=True):
        c1,c2,c3=st.columns([2,1,1])
        search=c1.text_input("Tìm mẫu, mẻ hoặc chỉ tiêu",placeholder="Nhập mã mẫu, tên mẻ, Benzen…")
        scope=c2.selectbox("Phạm vi",["Đang xử lý","Tất cả","Nhận hôm nay","Mẫu còn tồn"])
        statuses=c3.multiselect("Trạng thái",STATUSES)
        with st.expander("Bộ lọc bổ sung"):
            f1,f2,f3=st.columns(3)
            batches=f1.multiselect("Mẻ phân tích",sorted(df_current["Tên Mẻ"].dropna().astype(str).unique()))
            owners=f2.multiselect("Người giữ mẫu",sorted(df_current["Người Giữ"].dropna().astype(str).unique()))
            matrices=f3.multiselect("Nền mẫu",sorted(df_current["Nền Mẫu"].dropna().astype(str).unique()))
            
    mask=pd.Series(True,index=df_current.index)
    if scope=="Đang xử lý":mask &= active
    elif scope=="Nhận hôm nay":mask &= df_current["Ngày Nhận"].eq(today_date)
    elif scope=="Mẫu còn tồn":mask &= active & (df_current["Ngày Nhận"]<today_date)
    if search:
        query=lab_norm(search)
        mask &= df_current[["Mã Mẫu","Tên Mẻ","Chỉ Tiêu"]].fillna('').astype(str).apply(lambda row:query in lab_norm(' '.join(row)),axis=1)
    for col,values in [("Trạng Thái",statuses),("Tên Mẻ",batches),("Người Giữ",owners),("Nền Mẫu",matrices)]:
        if values:mask &= df_current[col].isin(values)
        
    df_display=df_current.loc[mask].sort_values("Giờ Nhận",ascending=False).copy()
    st.caption(f"Hiển thị {len(df_display)} / {len(df_current)} mẫu · Thay đổi chỉ được ghi khi bấm Lưu.")
    
    edited_df=st.data_editor(df_display,column_order=["Mã Mẫu","Tên Mẻ","Trạng Thái","Nền Mẫu","Chỉ Tiêu","Người Giữ","Ghi Chú","Giờ Nhận"],column_config={"Trạng Thái":st.column_config.SelectboxColumn(options=STATUSES,required=True),"Nền Mẫu":st.column_config.SelectboxColumn(options=["Khí","Nước","Chưa xác định"],required=True),"Giờ Nhận":st.column_config.DatetimeColumn(format="DD/MM/YYYY HH:mm"),"Ngày Nhận":None},disabled=["Mã Mẫu","Tên Mẻ","Chỉ Tiêu","Giờ Nhận","Ngày Nhận"],hide_index=True,use_container_width=True,num_rows="fixed",key="work_editor",height=350)
    
    savecol,exportcol=st.columns([1,3])
    if savecol.button("Lưu cập nhật",type="primary",use_container_width=True):
        candidate=st.session_state.df.copy()
        cols=["Trạng Thái","Nền Mẫu","Người Giữ","Ghi Chú"]
        candidate.loc[edited_df.index,cols]=edited_df[cols]
        try:
            save_data(candidate)
            st.session_state.df=candidate
            st.toast("Đã lưu cập nhật")
            st.rerun()
        except Exception as exc:
            st.error(f"Chưa lưu được: {exc}. Dữ liệu trong phiên chưa bị thay thế.")
    exportcol.download_button("↓ Xuất danh sách đang lọc",df_display.drop(columns=["Ngày Nhận"]).to_csv(index=False).encode('utf-8-sig'),"Danh_sach_mau.csv","text/csv")
    
    with st.expander("Phân công & chất lượng dữ liệu"):
        st.dataframe(df_current.loc[active].groupby("Người Giữ",dropna=False).size().reset_index(name="Mẫu đang xử lý"),hide_index=True,use_container_width=True)
        duplicated=df_current["Mã Mẫu"].fillna('').astype(str).str.strip().str.casefold().duplicated(keep=False)
        st.caption(f"{int(duplicated.sum())} dòng có mã mẫu trùng · {int(df_current['Giờ Nhận'].isna().sum())} dòng thiếu ngày nhận hợp lệ.")
        if duplicated.any():st.dataframe(df_current.loc[duplicated],hide_index=True)

elif menu == "📥 Quản lý Tiếp nhận":
    st.markdown("<h1 class='main-title'>📥 Khu vực Tiếp nhận mẫu mới</h1>", unsafe_allow_html=True)
    
    tab_excel, tab_thu_cong = st.tabs(["📁 Nạp qua Excel tự động", "✍️ Nhập Mẫu Lẻ"])
    
    with tab_excel:
        with st.container(border=True):
            st.info("💡 **Tính năng AI:** Tự động loại bỏ các ô bị bôi xám (#808080) (Các chỉ tiêu khách không yêu cầu).")
            uploaded_file = st.file_uploader("Kéo thả file KetQuaMeThuNghiem...xlsx vào đây", type=["xlsx", "xls"])
            
            if uploaded_file is not None:
                try:
                    import openpyxl
                    wb = openpyxl.load_workbook(io.BytesIO(uploaded_file.getvalue()), data_only=True)
                    ws = wb['Kết quả'] if 'Kết quả' in wb.sheetnames else wb.worksheets[0]
                    
                    batch, header_row, khm_col = "Không xác định", None, None
                    
                    for row in ws.iter_rows(min_row=1, max_row=min(30, ws.max_row)):
                        for cell in row:
                            val = clean_str(cell.value)
                            if val.startswith('Số:'):
                                batch = val.replace("Số:", "").strip()
                            if val.upper() == 'KHM':
                                header_row, khm_col = cell.row, cell.column
                                
                    if header_row is not None:
                        params_info = []
                        for col in range(khm_col + 1, ws.max_column + 1):
                            header_val = clean_str(ws.cell(header_row, col).value)
                            if header_val.casefold() == 'ghi chú' or not header_val: break
                            
                            header_val = re.sub(r'^\d+\.\s*', '', header_val)
                            header_val = re.sub(r'\s*\(chọn cái này\)', '', header_val, flags=re.I).strip()
                            params_info.append((col, header_val))
                        
                        samples_data = []
                        for r in range(header_row + 1, ws.max_row + 1):
                            khm_val = clean_str(ws.cell(r, khm_col).value)
                            if len(khm_val) < 3 or khm_val.lower() == 'nan': continue
                            
                            nen_mau_auto = "Khí" if khm_val.upper().startswith(("KT", "KKXQ", "KLV")) else ("Nước" if khm_val.upper().startswith(("NS", "NT", "NM", "NN")) else "Chưa xác định")

                            selected_params = []
                            for col_idx, param_name in params_info:
                                fill = ws.cell(r, col_idx).fill
                                color = fill.fgColor
                                
                                is_gray = fill.patternType == 'solid' and color.type == 'rgb' and str(color.rgb)[-6:].upper() == '808080'
                                if not is_gray: selected_params.append(param_name)
                                    
                            chuoi_chi_tieu = "; ".join(selected_params) if selected_params else "Chưa xác định"
                            samples_data.append({"Chọn": True, "Mã Mẫu": khm_val, "Nền Mẫu": nen_mau_auto, "Chỉ Tiêu": chuoi_chi_tieu})
                        
                        if len(samples_data) > 0:
                            st.success(f"✔️ Quét thành công **{len(samples_data)}** mẫu thuộc mẻ: **{batch}**")
                            edited_preview = st.data_editor(
                                pd.DataFrame(samples_data),
                                column_config={
                                    "Chọn": st.column_config.CheckboxColumn("Nhập mẫu?", default=True),
                                    "Mã Mẫu": st.column_config.TextColumn(disabled=True),
                                    "Nền Mẫu": st.column_config.SelectboxColumn("Nền Mẫu", options=["Nước", "Khí", "Chưa xác định"]),
                                    "Chỉ Tiêu": st.column_config.TextColumn(disabled=True)
                                },
                                hide_index=True, use_container_width=True
                            )
                            batch_nguoi = st.selectbox("Người tiếp nhận:", ["Thành", "Kỹ thuật viên 2", "Kỹ thuật viên 3"])
                            selected_samples = edited_preview[edited_preview["Chọn"] == True]
                            
                            if st.button(f"🚀 Lưu {len(selected_samples)} mẫu đã chọn vào Hệ thống", type="primary"):
                                new_rows = [{"Mã Mẫu": row["Mã Mẫu"], "Tên Mẻ": batch, "Nền Mẫu": row["Nền Mẫu"], "Chỉ Tiêu": row["Chỉ Tiêu"], "Trạng Thái": STATUSES[0], "Người Giữ": batch_nguoi, "Ghi Chú": "Import Excel", "Giờ Nhận": datetime.now()} for _, row in selected_samples.iterrows()]
                                st.session_state.df = pd.concat([st.session_state.df, pd.DataFrame(new_rows)], ignore_index=True)
                                save_data(st.session_state.df)
                                st.success("✅ Đã nạp thành công vào hệ thống!")
                                st.rerun()
                        else: st.warning("Không tìm thấy dữ liệu mẫu hợp lệ bên dưới ô KHM.")
                    else: st.error("Không tìm thấy ô 'KHM' trong file Excel!")
                except Exception as e: st.error(f"Lỗi đọc file: {e}")

    with tab_thu_cong:
        with st.container(border=True):
            with st.form("add_sample_form", clear_on_submit=True):
                st.markdown("<div class='sub-title'>Bổ sung Mẫu Lẻ</div>", unsafe_allow_html=True)
                new_id = st.text_input("Mã Mẫu (VD: NS-1509-01)*")
                new_name = st.text_input("Tên Mẻ (VD: 2026.07.017)")
                col_t1, col_t2 = st.columns(2)
                with col_t1: new_nen = st.selectbox("Nền Mẫu", ["Nước", "Khí"])
                with col_t2: new_chi_tieu = st.text_input("Chỉ tiêu đo")
                new_nguoi = st.text_input("Người tiếp nhận (Ký tên)")
                
                if st.form_submit_button("Thêm Mẫu lẻ", type="primary") and new_id:
                    st.session_state.df = pd.concat([st.session_state.df, pd.DataFrame([{"Mã Mẫu": new_id, "Tên Mẻ": new_name, "Nền Mẫu": new_nen, "Chỉ Tiêu": new_chi_tieu, "Trạng Thái": STATUSES[0], "Người Giữ": new_nguoi, "Ghi Chú": "", "Giờ Nhận": datetime.now()}])], ignore_index=True)
                    save_data(st.session_state.df)
                    st.success(f"✅ Đã thêm mẫu {new_id} thành công!")

elif menu == "⚙️ Vận hành GC-MS":
    st.markdown("<h1 class='main-title'>⚙️ Phân tích & Vận hành Máy đo GC-MS</h1>", unsafe_allow_html=True)
    
    tab_seq, tab_auto = st.tabs(["1. Xuất Sequence Chạy Máy", "2. Xử lý Tự động (File PDF/Excel)"])
    
    with tab_seq:
        with st.container(border=True):
            st.markdown("<div class='sub-title'>Xuất Sequence Chạy Máy</div>", unsafe_allow_html=True)
            df_ready = st.session_state.df[st.session_state.df["Trạng Thái"] == "🟡 3. Chờ chạy máy"]
            st.write(f"Đang có **{len(df_ready)}** mẫu trong hàng chờ.")
            
            if not df_ready.empty:
                seq_df = pd.DataFrame({'Vial': range(1, len(df_ready) + 1), 'Sample Name': df_ready['Mã Mẫu'], 'Sample Type': 'Sample'})
                seq_df['Method'] = df_ready['Chỉ Tiêu'].apply(lambda x: 'VOCs.M' if any(k in str(x).upper() for k in ['VOC', 'BENZEN', 'TOLUEN', 'CHLORO', 'STYREN']) else 'HCHO.M')
                seq_df['Data File'] = datetime.now().strftime("%Y%m%d") + "_" + df_ready['Mã Mẫu']
                st.download_button("📥 Tải File Sequence.csv", data=seq_df.to_csv(index=False).encode('utf-8'), file_name=f"MassHunter_Seq_{datetime.now().strftime('%Y%m%d')}.csv", mime="text/csv", type="primary")
            
    with tab_auto:
        unit_cfg=u_controls("auto_")
        with st.container(border=True):
            st.markdown("<div class='sub-title'>Xử lý Kết quả Hàng loạt (SOP)</div>", unsafe_allow_html=True)
            st.info("💡 Tự động bóc tách số liệu, nội suy nồng độ $C_{surr}$ chuẩn và so khớp Giới hạn MDL/LOQ theo đúng chuẩn phòng Lab.")
            
            gc_file = st.file_uploader("Kéo thả báo cáo GC (PDF/Excel/CSV)", type=["pdf", "xlsx", "xls", "csv"])

            if gc_file is not None and unit_cfg["confirmed"]:
                calc_results = []
                try:
                    dynamic_compounds = []
                    if not st.session_state.df_limit.empty and "Tên Chất" in st.session_state.df_limit.columns:
                        dynamic_compounds = st.session_state.df_limit["Tên Chất"].dropna().astype(str).str.strip().tolist()
                    
                    surrogate_compounds = ['Toluene-D8', 'Toluen-D8', 'BFB', '4-Bromofluorobenzene', 'Chlorobenzene-d5']
                    known_compounds_upper = set(c.upper() for c in dynamic_compounds + surrogate_compounds)

                    if gc_file.name.endswith('.pdf'):
                        import PyPDF2
                        text = "".join([page.extract_text() + "\n" for page in PyPDF2.PdfReader(gc_file).pages])
                        pdf_data, current_compound = [], None

                        for line in text.split('\n'):
                            parts = line.split()
                            if not parts: continue
                            
                            clean_line = line.strip()
                            if clean_line.upper() in known_compounds_upper: 
                                current_compound = clean_line
                                continue
                            
                            data_file_raw = parts[0].replace('.d', '')
                            if len(parts) >= 5 and (parts[0].endswith('.d') or data_file_raw in ['10PPM', '10a', '2', '4', '6', '8']) and ('Sample' in parts or 'Cal' in parts):
                                data_file = data_file_raw
                                floats = [float(p.replace(',', '.')) for p in parts if p.replace('.', '', 1).replace(',', '', 1).isdigit() and (p.count('.') + p.count(',') <= 1)]
                                
                                if len(floats) >= 3 and current_compound:
                                    final_conc = floats[-3] if 'Cal' in parts and len(floats) >= 5 else (floats[-2] if 'Cal' in parts and len(floats) >= 4 else floats[-1])
                                    pdf_data.append({"Data File": data_file, "Compound Name": current_compound, "Final Conc.": final_conc})
                        
                        df_gc, compound_col = pd.DataFrame(pdf_data), "Compound Name"
                    else:
                        df_gc = pd.read_csv(gc_file) if gc_file.name.endswith('.csv') else pd.read_excel(gc_file)
                        df_gc.columns = [str(c).strip() for c in df_gc.columns]
                        compound_col = next((c for c in df_gc.columns if c.lower() in ['name', 'compound', 'compound name', 'tên chất']), None)

                    if 'Data File' in df_gc.columns and 'Final Conc.' in df_gc.columns and compound_col:
                        unit_column=next((c for c in ['Units','Unit','Đơn vị'] if c in df_gc.columns),None)
                        if unit_column:
                            declared={u_canonical(v) for v in df_gc[unit_column].dropna() if str(v).strip()}
                            if declared and declared!={u_canonical(unit_cfg['src'])}:raise ValueError('Đơn vị trong file khác đơn vị đã chọn hoặc có nhiều đơn vị. Chọn đúng đơn vị nguồn hoặc tách file theo đơn vị trước khi tính.')
                        surrogate_dict = {}
                        for _, row in df_gc.iterrows():
                            comp_name = str(row[compound_col]).upper()
                            if comp_name in ['TOLUENE-D8', 'TOLUEN-D8', 'BFB', '4-BROMOFLUOROBENZENE']:
                                sample_name = str(row['Data File']).replace('.d', '')
                                raw_conc = pd.to_numeric(row['Final Conc.'], errors='coerce')
                                if pd.notna(raw_conc) and raw_conc > 0:
                                    raw_conc = u_convert(raw_conc,unit_cfg["src"],unit_cfg["ref"])
                                    c_exp = get_dynamic_surrogate_expected(raw_conc)
                                    recovery = (raw_conc / c_exp) * 100.0
                                    surrogate_dict[sample_name] = {"recovery": recovery, "c_exp": c_exp, "c_do": raw_conc}

                        default_v_gas, default_nen, default_loai = 24.0, 'KT', 'Khí'
                        for _, row in df_gc.iterrows():
                            sn = str(row['Data File']).upper()
                            if 'KT' in sn: default_v_gas, default_nen, default_loai = 24.0, 'KT', 'Khí'; break
                            elif 'KXQ' in sn: default_v_gas, default_nen, default_loai = 4.0, 'KXQ', 'Khí'; break
                            elif 'KLV' in sn: default_v_gas, default_nen, default_loai = 4.0, 'KLV', 'Khí'; break
                            elif any(k in sn for k in ['NS', 'NT', 'NM', 'NN']): default_v_gas, default_nen, default_loai = 1.0, 'NS', 'Nước'; break

                        for _, row in df_gc.iterrows():
                            sample_name, comp_name = str(row['Data File']).replace('.d', ''), str(row[compound_col])
                            upper_name = sample_name.upper()
                            
                            if not any(k in upper_name for k in ['KT', 'KXQ', 'KLV', 'NS', 'NT', 'NM', 'NN', 'BL', 'BLANK', 'TC', 'QC']): continue
                            if upper_name in ['1', '2', '4', '5', '6', '8', '10', '10A'] or 'PPM' in upper_name: continue
                            if comp_name.upper() in ['TOLUENE-D8', 'TOLUEN-D8', 'BFB', '4-BROMOFLUOROBENZENE']: continue
                                
                            raw_conc = pd.to_numeric(row['Final Conc.'], errors='coerce')
                            if pd.isna(raw_conc): continue

                            v_param, nen_mau, loai_mau = parse_sample_matrix(sample_name)
                            if v_param is None: v_param, nen_mau, loai_mau = default_v_gas, default_nen, default_loai
                                
                            surr_info = surrogate_dict.get(sample_name, {"recovery": 100.0, "c_exp": "", "c_do": ""})
                            sample_recovery = surr_info["recovery"]
                            c_surr_truoc = surr_info["c_exp"]
                            c_surr_sau = surr_info["c_do"]
                            
                            mdl_val, loq_val, unit = get_limit_info(comp_name, nen_mau)
                            
                            c_thuc_str = evaluate_result(raw_conc, v_param, mdl_val, loq_val, unit, loai_mau, recovery=sample_recovery)
                            
                            limit_display = ""
                            if loai_mau == 'Khí' and mdl_val is not None: limit_display = f"MDL: {mdl_val} {unit}"
                            elif loai_mau == 'Nước' and loq_val is not None: limit_display = f"LOQ: {loq_val} {unit}"
                            
                            calc_results.append({
                                "Tên mẫu": sample_name, "Tên chỉ tiêu": comp_name, "C đo": round(raw_conc, 4), 
                                "C thực": c_thuc_str, "Giới hạn": limit_display, "R(%)": f"{round(sample_recovery, 1)}%",
                                "C_surr_truoc": c_surr_truoc, "C_surr_sau": c_surr_sau,
                                "Đơn vị C đo":unit_cfg["src"],"Đơn vị Csurr":unit_cfg["ref"],"Đơn vị kết quả":unit_cfg["water"] if loai_mau=="Nước" else unit_cfg["gas_out"]
                            })

                    if calc_results:
                        st.success(f"✅ Đã xử lý {len(calc_results)} dòng kết quả. Tự động áp dụng tiêu chuẩn Khí/Nước.")
                        
                        df_results = pd.DataFrame(calc_results)
                        df_results = df_results.sort_values(by=["Tên mẫu", "Tên chỉ tiêu"]).reset_index(drop=True)
                        
                        st.session_state.results = df_results
                        st.session_state.results_stale = False
                        
                        display_cols = ["Tên mẫu", "Tên chỉ tiêu", "C đo", "Đơn vị C đo", "C thực", "Đơn vị kết quả", "Giới hạn", "R(%)"]
                        st.dataframe(df_results[display_cols], use_container_width=True, hide_index=True)
                        
                        csv_results = df_results.to_csv(index=False).encode('utf-8-sig')
                        st.download_button(
                            label="📥 Tải Kết quả (CSV) để Lưu trữ",
                            data=csv_results,
                            file_name=f"Ket_Qua_GC_{datetime.now().strftime('%Y%m%d_%H%M')}.csv",
                            mime="text/csv",
                            type="primary"
                        )
                        
                        processed_samples = df_results["Tên mẫu"].unique().tolist()
                        for smp in processed_samples:
                            mask = st.session_state.df["Mã Mẫu"] == smp
                            if mask.any():
                                st.session_state.df.loc[mask, "Trạng Thái"] = "🟣 5. Đang tính số liệu"
                        save_data(st.session_state.df)
                        
                    else: 
                        st.warning("⚠️ Báo cáo không chứa mẫu hợp lệ (KT, KXQ, NS, NT...) hoặc thiếu dữ liệu phân tích.")
                except Exception as e: st.error(f"❌ Lỗi xử lý: {e}")

elif menu == "🔥 Vận hành GC-FID":
    st.markdown("<h1 class='main-title'>🔥 Hệ thống GC-FID (Agilent)</h1>", unsafe_allow_html=True)
    st.caption("Module chuyên biệt xử lý dữ liệu từ đầu dò FID")
    
    col_seq, col_import = st.columns(2)
    with col_seq:
        with st.container(border=True):
            st.markdown("<div class='sub-title'>1. Xuất Sequence GC-FID</div>", unsafe_allow_html=True)
            st.info("Sẽ tích hợp thuật toán xuất file Sequence định dạng cho máy GC-FID Agilent.")
        
    with col_import:
        with st.container(border=True):
            st.markdown("<div class='sub-title'>2. Xử lý kết quả GC-FID</div>", unsafe_allow_html=True)
            st.info("Khu vực chờ tích hợp thuật toán đọc file báo cáo từ máy GC-FID.")
            fid_file = st.file_uploader("Kéo thả báo cáo GC-FID (PDF/Excel/CSV/TXT)", type=["pdf", "xlsx", "xls", "csv", "txt"])
            if fid_file:
                st.warning("🚧 Hệ thống đang chờ cập nhật thuật toán bóc tách dữ liệu từ file report FID.")

elif menu == "🧬 Vận hành Thermo":
    st.markdown("<h1 class='main-title'>🧬 Hệ thống Thermo GC-MS</h1>", unsafe_allow_html=True)
    st.caption("Module chuyên biệt xử lý dữ liệu OCP, OPP, PCB và Phenol")
    
    col_seq, col_import = st.columns(2)
    with col_seq:
        with st.container(border=True):
            st.markdown("<div class='sub-title'>1. Xuất Sequence Thermo</div>", unsafe_allow_html=True)
            st.info("Sẽ tích hợp thuật toán xuất file Sequence định dạng tương thích phần mềm Thermo.")
        
    with col_import:
        with st.container(border=True):
            st.markdown("<div class='sub-title'>2. Xử lý kết quả Thermo</div>", unsafe_allow_html=True)
            st.info("Khu vực chờ tích hợp thuật toán đọc file xuất từ máy Thermo.")
            thermo_file = st.file_uploader("Kéo thả báo cáo Thermo (PDF/Excel/CSV)", type=["pdf", "xlsx", "xls", "csv"])
            if thermo_file:
                st.warning("🚧 Hệ thống đang chờ cập nhật thuật toán bóc tách dữ liệu từ file report Thermo.")

elif menu == "🧮 Tiện ích Phân tích":
    st.markdown("<h1 class='main-title'>🧮 Tiện ích Phân tích & Tính toán</h1>", unsafe_allow_html=True)
    
    tab_manual, tab_calib = st.tabs(["1. Tính toán & Nhập liệu Thủ công", "2. Pha Đường Chuẩn Đa Thành Phần"])
    
    with tab_manual:
        unit_cfg=u_controls("manual_")
        with st.container(border=True):
            st.markdown("<div class='sub-title'>Tính toán & Nhập liệu Thủ công (Gồm tính Tổng)</div>", unsafe_allow_html=True)
            st.info("💡 Điền thông số đo để máy tự tính, hoặc nhập thẳng vào cột 'Kết quả'. Có thể thêm dòng chỉ tiêu bị thiếu. Đánh dấu các chất cần tính dồn để cộng thành một chỉ tiêu Tổng chung.")
            
            valid_samples = df_current[df_current['Chỉ Tiêu'].str.strip() != ""]['Mã Mẫu'].tolist()
            selected_sample_manual = st.selectbox("🔍 Chọn Mã Mẫu để nhập liệu:", ["-- Chọn mẫu --"] + valid_samples)
            
            if selected_sample_manual != "-- Chọn mẫu --":
                sample_info = df_current[df_current['Mã Mẫu'] == selected_sample_manual].iloc[0]
                chi_tieu_raw = str(sample_info['Chỉ Tiêu'])
                chi_tieu_list = [ct.strip() for ct in re.split(r'[;,]', chi_tieu_raw) if ct.strip()]
                
                v_param_def, nen_mau_code, loai_mau_def = parse_sample_matrix(selected_sample_manual)
                if v_param_def is None:
                    v_param_def, nen_mau_code, loai_mau_def = 24.0, 'KT', 'Khí'
                
                st.write(f"**Phân loại:** {loai_mau_def} ({nen_mau_code}) | **Chỉ tiêu theo Database:** {len(chi_tieu_list)}")
                
                manual_data = []
                for ct in chi_tieu_list:
                    mdl_val, loq_val, unit = get_limit_info(ct, nen_mau_code)
                    limit_str = ""
                    if loai_mau_def == 'Khí' and mdl_val is not None: limit_str = f"MDL: {mdl_val} {unit}"
                    elif loai_mau_def == 'Nước' and loq_val is not None: limit_str = f"LOQ: {loq_val} {unit}"

                    manual_data.append({
                        "Cộng Tổng": False,
                        "Chỉ Tiêu": ct,
                        "Giới hạn (Tới hạn)": limit_str,
                        "Độ làm giàu (V)": v_param_def,
                        "C_surr thực": 10.0,
                        "C_surr đo": 10.0,
                        "C_đo chỉ tiêu": 0.0,
                        "Kết quả (Ghi đè)": ""
                    })
                
                df_manual = pd.DataFrame(manual_data)
                
                st.caption("Có thể thêm hàng mới (dấu + ở dưới bảng) nếu mẫu có chỉ tiêu con bị sót.")
                edited_manual = st.data_editor(
                    df_manual,
                    num_rows="dynamic",
                    column_config={
                        "Cộng Tổng": st.column_config.CheckboxColumn("Cộng Tổng?"),
                        "Chỉ Tiêu": st.column_config.TextColumn("Chỉ Tiêu"),
                        "Giới hạn (Tới hạn)": st.column_config.TextColumn(disabled=True),
                        "Độ làm giàu (V)": st.column_config.NumberColumn(format="%.2f"),
                        "C_surr thực": st.column_config.NumberColumn(format="%.3f"),
                        "C_surr đo": st.column_config.NumberColumn(format="%.3f"),
                        "C_đo chỉ tiêu": st.column_config.NumberColumn(format="%.4f"),
                        "Kết quả (Ghi đè)": st.column_config.TextColumn("Kết quả (Tự nhập)")
                    },
                    use_container_width=True,
                    hide_index=True,
                    key=f"manual_editor_{selected_sample_manual}"
                )
                
                st.markdown("---")
                total_param_name = st.text_input("📝 Nhập tên chỉ tiêu Tổng (Chỉ áp dụng nếu có tick chọn 'Cộng Tổng' ở bảng trên):", placeholder="VD: Tổng VOCs, Tổng PCB...")
                
                if st.button("💾 Tính Toán & Lưu Mẫu Này", type="primary",disabled=not unit_cfg["confirmed"]):
                    try:
                        new_results = []
                        sum_val = 0.0
                        has_sum = False
                    
                        for _, row in edited_manual.iterrows():
                            ct = str(row["Chỉ Tiêu"]).strip()
                            if not ct or ct == "nan": continue
                        
                            v = float(row.get("Độ làm giàu (V)", v_param_def) or v_param_def)
                            c_truoc = float(row.get("C_surr thực", 10.0) or 10.0)
                            c_sau = float(row.get("C_surr đo", 10.0) or 10.0)
                            c_do = float(row.get("C_đo chỉ tiêu", 0.0) or 0.0)
                            override_res = str(row.get("Kết quả (Ghi đè)", "")).strip()
                            is_sum = bool(row.get("Cộng Tổng", False))
                        
                            c_sau=u_convert(c_sau,unit_cfg["src"],unit_cfg["ref"])
                            if c_truoc<=0: raise ValueError("Csurr trước phải >0")
                            recovery = (c_sau / c_truoc) * 100.0
                            mdl_val, loq_val, unit = get_limit_info(ct, nen_mau_code)
                        
                            limit_display = ""
                            if loai_mau_def == 'Khí' and mdl_val is not None: limit_display = f"MDL: {mdl_val} {unit}"
                            elif loai_mau_def == 'Nước' and loq_val is not None: limit_display = f"LOQ: {loq_val} {unit}"
                        
                            c_thuc_str = ""
                            c_thuc_numeric = 0.0
                        
                            if override_res and override_res.lower() != "nan":
                                c_thuc_str = override_res
                                try: c_thuc_numeric = float(override_res.replace(',', '.'))
                                except: c_thuc_numeric = 0.0
                            else:
                                c_thuc_str = evaluate_result(c_do, v, mdl_val, loq_val, unit, loai_mau_def, recovery)
                                if c_thuc_str.startswith("KPH") or c_thuc_str.startswith("<"):
                                    c_thuc_numeric = 0.0
                                else:
                                    try: c_thuc_numeric = float(c_thuc_str.replace(',', '.'))
                                    except: c_thuc_numeric = 0.0
                        
                            if is_sum:
                                has_sum = True
                                sum_val += c_thuc_numeric
                            
                            new_results.append({
                                "Tên mẫu": selected_sample_manual,
                                "Tên chỉ tiêu": ct,
                                "C đo": round(c_do, 4),
                                "C thực": c_thuc_str,
                                "Giới hạn": limit_display,
                                "R(%)": f"{round(recovery, 1)}%",
                                "C_surr_truoc": c_truoc,
                                "C_surr_sau": c_sau,
                                "Đơn vị C đo":unit_cfg["src"],"Đơn vị Csurr":unit_cfg["ref"],"Đơn vị kết quả":unit_cfg["water"] if loai_mau_def=="Nước" else unit_cfg["gas_out"]
                            })
                    
                        if has_sum and total_param_name:
                            new_results.append({
                                "Tên mẫu": selected_sample_manual,
                                "Tên chỉ tiêu": total_param_name,
                                "C đo": "",
                                "C thực": str(round(sum_val, 4)).replace('.', ','),
                                "Giới hạn": "", 
                                "R(%)": "",
                                "C_surr_truoc": "",
                                "C_surr_sau": "",
                                "Đơn vị kết quả":unit_cfg["water"] if loai_mau_def=="Nước" else unit_cfg["gas_out"]
                            })
                    
                        new_df = pd.DataFrame(new_results)
                        if 'results' not in st.session_state or st.session_state.results.empty:
                            st.session_state.results = new_df
                        else:
                            existing = st.session_state.results
                            existing = existing[existing["Tên mẫu"] != selected_sample_manual]
                            st.session_state.results = pd.concat([existing, new_df], ignore_index=True)
                    
                        st.session_state.results_stale = False
                    
                        mask = st.session_state.df["Mã Mẫu"] == selected_sample_manual
                        if mask.any():
                            st.session_state.df.loc[mask, "Trạng Thái"] = "🟣 5. Đang tính số liệu"
                            save_data(st.session_state.df)
                        
                        st.success(f"✅ Đã lưu kết quả cho mẫu {selected_sample_manual}! Hệ thống đã cộng dồn vào danh sách tổng chờ xuất Biên bản.")
                        st.rerun()
                    except Exception as exc:
                        st.error(str(exc))

    with tab_calib:
        calib_v_unit=st.selectbox("Đơn vị thể tích định mức",["mL","µL","L"],key="calib_v_unit")
        calib_unit=st.selectbox("Đơn vị nhập dãy chuẩn và nồng độ đích",list(U_GROUPS["Nồng độ dung dịch"]),index=5,key="calib_unit")
        st.caption("Chọn Đơn vị gốc cho từng dòng; kết quả công thức tự quy về µg/mL. Các thể tích hút xuất bằng µL.")
        with st.container(border=True):
            st.markdown("<div class='sub-title'>🧪 Lập Công Thức Pha Chuẩn Tối Ưu (Smart Mix)</div>", unsafe_allow_html=True)
            st.info("💡 Hệ thống hỗ trợ tối ưu hóa quy trình pha chuẩn: Gộp các Nội chuẩn/Surrogate thành 1 Master Mix để hút 1 lần; Gộp các Mix chuẩn gốc thành Mix làm việc (Working Standard) để tránh sai số pipet nhỏ.")
            
            col_v1, col_v2 = st.columns(2)
            with col_v1:
                v_final = st.number_input(f"Thể tích định mức mỗi điểm chuẩn ({calib_v_unit})", value=1.0, step=0.1, min_value=0.1)
            with col_v2:
                levels_input = st.text_input(f"🎯 Dãy chuẩn ({calib_unit}), phân cách bằng dấu phẩy; thập phân dùng dấu chấm:", "0.5, 1, 2, 5, 10, 20")
            
            st.markdown("### ⚙️ Tùy chọn Tối ưu hóa (Optimization)")
            col_opt1, col_opt2 = st.columns(2)
            with col_opt1:
                opt_mix_is_surr = st.toggle("🧪 Gộp IS & Surrogate thành Master Mix", value=True)
                if opt_mix_is_surr:
                    v_spike_is = st.number_input("Thể tích hút Master Mix cho mỗi vial (µL)", value=50.0, step=10.0)
            with col_opt2:
                opt_mix_std = st.toggle("🧪 Gộp các Chuẩn gốc thành Mix Trung gian", value=True)
                if opt_mix_std:
                    c_ws_std = st.number_input(f"Nồng độ Mix Trung gian ({calib_unit})", value=100.0, step=10.0)
                    v_ws_total = st.number_input("Thể tích cần pha Mix Trung gian (µL)", value=1000.0, step=100.0)

            st.markdown("**1. Các dung dịch Chuẩn (Mix) thay đổi theo dãy nồng độ:**")
            if "df_mix_default" not in st.session_state:
                st.session_state.df_mix_default = pd.DataFrame([{"Tên Mix Chuẩn": "Mix VOCs", "C gốc (giá trị nhập)": 1000.0, "Đơn vị gốc":"µg/mL"}])
            df_mix = st.data_editor(st.session_state.df_mix_default, num_rows="dynamic", use_container_width=True, key="mix_editor")
            
            st.markdown("**2. Các dung dịch Nội chuẩn (IS) & Đồng hành (Surrogate) cố định:**")
            if "df_is_default" not in st.session_state:
                st.session_state.df_is_default = pd.DataFrame([
                    {"Phân Loại": "Nội chuẩn (IS)", "Tên Hợp Chất": "Fluorobenzene", "C gốc (giá trị nhập)": 1000.0, "Đơn vị gốc":"µg/mL", "C đích (đơn vị đang chọn)": 10.0},
                    {"Phân Loại": "Surrogate", "Tên Hợp Chất": "Toluene-D8", "C gốc (giá trị nhập)": 1000.0, "Đơn vị gốc":"µg/mL", "C đích (đơn vị đang chọn)": 10.0}
                ])
            df_is_surr = st.data_editor(st.session_state.df_is_default, num_rows="dynamic", use_container_width=True, key="is_surr_editor", column_config={"Phân Loại": st.column_config.SelectboxColumn(options=["Nội chuẩn (IS)", "Surrogate", "Khác"])})
            
            if st.button("🚀 Tính Toán Bảng Pha Chuẩn & Tối Ưu", type="primary"):
                try:
                    v_final=u_convert(v_final,calib_v_unit,"mL")
                    levels = sorted([u_convert(float(x.strip()),calib_unit,"µg/mL") for x in levels_input.split(",") if x.strip()])
                    if not levels or any(x<0 for x in levels):raise ValueError("Dãy chuẩn không hợp lệ")
                    if opt_mix_std:c_ws_std=u_convert(c_ws_std,calib_unit,"µg/mL")
                    if opt_mix_is_surr and v_spike_is<=0:raise ValueError("Thể tích hút phải >0")
                    for frame in [df_mix,df_is_surr]:
                        for _,r in frame.iterrows():
                            if float(r.get("C gốc (giá trị nhập)",0))<=0:raise ValueError("Nồng độ gốc phải >0")
                    if opt_mix_std and (c_ws_std<=0 or v_ws_total<=0):raise ValueError("Thông số mix trung gian phải >0")
                    calib_data = []
                    instructions = []

                    is_surr_columns = {} 
                    total_is_v_per_vial = 0.0
                    
                    if opt_mix_is_surr:
                        v_master_total = (len(levels) + 3) * v_spike_is
                        instructions.append("### 🧪 BƯỚC 1: PHA MASTER MIX NỘI chuẩn (IS) & SURROGATE")
                        instructions.append(f"*(Pha tổng cộng {v_master_total} µL Master Mix dùng chung cho toàn bộ các điểm chuẩn. Mỗi điểm chuẩn sẽ hút {v_spike_is} µL)*")
                        sum_v = 0.0
                        for _, row in df_is_surr.iterrows():
                            name = str(row.get("Tên Hợp Chất", "")).strip()
                            if not name or name == "nan": continue
                            c_s = u_convert(float(row.get("C gốc (giá trị nhập)", 0)),row.get("Đơn vị gốc","µg/mL"),"µg/mL")
                            if c_s<=0:raise ValueError("Nồng độ chuẩn gốc phải >0")
                            c_t = u_convert(float(row.get("C đích (đơn vị đang chọn)", 0)),calib_unit,"µg/mL")
                            v_stock = (c_t * v_final * 1000 * v_master_total) / (c_s * v_spike_is) if c_s > 0 else 0
                            instructions.append(f"- Hút **{v_stock:.2f} µL** {name} gốc ({c_s} µg/mL)")
                            sum_v += v_stock
                        v_solv = v_master_total - sum_v
                        instructions.append(f"- Thêm **{v_solv:.2f} µL** dung môi. Lắc đều.")
                        if sum_v > v_master_total:raise ValueError("Master Mix vượt thể tích. Giảm nồng độ đích hoặc tăng thể tích hút.")
                        
                        is_surr_columns["Hút IS/Surr Master Mix (µL)"] = v_spike_is
                        total_is_v_per_vial = v_spike_is
                    else:
                        for _, row in df_is_surr.iterrows():
                            name = str(row.get("Tên Hợp Chất", "")).strip()
                            if not name or name == "nan": continue
                            c_s = u_convert(float(row.get("C gốc (giá trị nhập)", 0)),row.get("Đơn vị gốc","µg/mL"),"µg/mL")
                            if c_s<=0:raise ValueError("Nồng độ chuẩn gốc phải >0")
                            c_t = u_convert(float(row.get("C đích (đơn vị đang chọn)", 0)),calib_unit,"µg/mL")
                            v_ul = (c_t * v_final * 1000) / c_s if c_s > 0 else 0
                            prefix = "IS" if "IS" in str(row.get("Phân Loại", "")) else "Surr"
                            is_surr_columns[f"Hút {prefix}: {name} (µL)"] = v_ul
                            total_is_v_per_vial += v_ul

                    if opt_mix_std:
                        instructions.append("### 🧪 BƯỚC 2: PHA MIX CHUẨN LÀM VIỆC (WORKING STANDARD)")
                        instructions.append(f"*(Gộp các Mix chuẩn gốc thành Mix trung gian duy nhất có nồng độ {c_ws_std} µg/mL. Thể tích pha: {v_ws_total} µL)*")
                        sum_v = 0.0
                        for _, r in df_mix.iterrows():
                            name = str(r.get("Tên Mix Chuẩn", "")).strip()
                            if not name or name == "nan": continue
                            c_s = u_convert(float(r.get("C gốc (giá trị nhập)", 0)),r.get("Đơn vị gốc","µg/mL"),"µg/mL")
                            v_stock = (c_ws_std * v_ws_total) / c_s if c_s > 0 else 0
                            instructions.append(f"- Hút **{v_stock:.2f} µL** {name} ({c_s} µg/mL)")
                            sum_v += v_stock
                        v_solv = v_ws_total - sum_v
                        instructions.append(f"- Thêm **{v_solv:.2f} µL** dung môi. Lắc đều.")
                        if sum_v > v_ws_total:raise ValueError("Mix trung gian vượt thể tích.")

                    instructions.append("### 🧪 BƯỚC 3: PHA DÃY CHUẨN VÀO VIAL CUỐI CÙNG")
                    for lvl in levels:
                        row_data = {"Điểm chuẩn": f"Level {lvl} µg/mL"}
                        total_std_v = 0.0
                        
                        if opt_mix_std:
                            v_ws = (lvl * v_final * 1000) / c_ws_std if c_ws_std > 0 else 0
                            row_data[f"Hút Mix Làm Việc {c_ws_std}µg/mL (µL)"] = round(v_ws, 2)
                            total_std_v += v_ws
                        else:
                            for _, r in df_mix.iterrows():
                                name = str(r.get("Tên Mix Chuẩn", "")).strip()
                                if not name or name == "nan": continue
                                c_s = u_convert(float(r.get("C gốc (giá trị nhập)", 0)),r.get("Đơn vị gốc","µg/mL"),"µg/mL")
                                v_ul = (lvl * v_final * 1000) / c_s if c_s > 0 else 0
                                row_data[f"Hút Mix {name} (µL)"] = round(v_ul, 2)
                                total_std_v += v_ul
                        
                        for k, v in is_surr_columns.items():
                            row_data[k] = round(v, 2)
                            
                        v_dungmoi = (v_final * 1000) - total_std_v - total_is_v_per_vial
                        if v_dungmoi<0:raise ValueError("Điểm chuẩn vượt thể tích vial; chưa xuất bảng pha.")
                        row_data["Dung môi bù (µL)"] = round(v_dungmoi, 2) if v_dungmoi >= 0 else "Quá thể tích!"
                        row_data["V tổng đích (mL)"] = v_final
                        calib_data.append(row_data)

                    for ins in instructions:
                        st.markdown(ins)
                    
                    df_calib = pd.DataFrame(calib_data)
                    st.dataframe(df_calib, use_container_width=True, hide_index=True)
                    
                    st.download_button("📥 Tải Bảng Pha Chuẩn (CSV)", data=df_calib.to_csv(index=False).encode('utf-8-sig'), file_name=f"Quy_trinh_Pha_chuan_{datetime.now().strftime('%Y%m%d')}.csv", mime="text/csv")
                except Exception as e:
                    st.error(f"Lỗi tính toán: {e}. Vui lòng kiểm tra lại dữ liệu.")

elif menu == "📝 Báo cáo & Lập Biên bản":
    st.markdown("<h1 class='main-title'>📝 Báo cáo & Lập Biên Bản</h1>", unsafe_allow_html=True)
    
    tab_report, tab_qr = st.tabs(["1. Lập Biên Bản Xử Lý Mẫu", "2. Sinh Mã Vạch QR"])

    with tab_report:
        st.info("Biên bản dùng thẻ {{ row.don_vi_cdo }}, {{ row.don_vi_csurr }}, {{ row.don_vi_kq }} để hiển thị đơn vị đã chọn. Template có đơn vị cố định cần sửa nhãn cho khớp; không tự thay toàn bộ chữ ppm trong mẫu.") 
        st.info("💡 **Hệ thống Thông minh:** Tự động lấy kết quả bạn vừa tính ở tab Phân tích để điền vào Biểu mẫu Word/Excel. Mọi định dạng Form, Chữ ký được bảo lưu 100%.")
        
        with st.container(border=True):
            col_tpl, col_data = st.columns(2)
            with col_tpl:
                template_file = st.file_uploader("1. Tải lên Form Mẫu (.docx, .xlsx)", type=["docx", "xlsx"])
            with col_data:
                data_file = st.file_uploader("2. File Số liệu (.csv) [Tùy chọn tải tay]", type=["xlsx", "csv"])

            if template_file:
                try:
                    if data_file is not None:
                        df_kq = pd.read_excel(data_file) if data_file.name.endswith('.xlsx') else pd.read_csv(data_file)
                    else:
                        df_kq = st.session_state.get('results', pd.DataFrame())
                    
                    if df_kq.empty or "Tên mẫu" not in df_kq.columns or "Tên chỉ tiêu" not in df_kq.columns:
                        st.error("⚠️ Hệ thống đang không có số liệu kết quả trong bộ nhớ. Bạn hãy qua tab **Vận hành Máy Đo** hoặc **Tiện ích Phân tích** tính số liệu trước, hoặc thả thủ công file `.csv` kết quả vào ô số 2 nhé!")
                    elif st.session_state.get('results_stale') and data_file is None:
                        st.warning("⚠️ **Cảnh báo lệch số liệu:** Thư viện giới hạn MDL/LOQ vừa bị sửa. Bạn nên quay lại tab phân tích bấm 'Tính lại' để kết quả được cập nhật chuẩn xác nhất!")
                    else:
                        ket_qua_list = []
                        danh_sach_mau = []
                        grouped = df_kq.groupby("Tên mẫu")
                        
                        for ten_mau, group in grouped:
                            first_row = group.iloc[0]
                            ten_mau_str = str(ten_mau)
                            
                            v_khi = "24,0" if "KT" in ten_mau_str.upper() else ("4,0" if any(k in ten_mau_str.upper() for k in ["KXQ", "KLV"]) else "")
                            h_phantram = str(first_row.get("R(%)", "")).replace('%', '').strip()
                            c_surr_truoc = str(first_row.get("C_surr_truoc", "")).replace('.', ',')
                            c_surr_sau = str(first_row.get("C_surr_sau", "")).replace('.', ',')
                            
                            if c_surr_sau == "" and h_phantram != "":
                                try: c_surr_sau = str(round((float(h_phantram) / 100) * 10.0, 2)).replace('.', ',')
                                except: pass

                            mau_dict = {
                                "ngay": datetime.now().strftime("%d/%m/%Y"),
                                "ky_hieu": ten_mau_str,
                                "c_surr_truoc": c_surr_truoc, 
                                "c_surr_sau": c_surr_sau, 
                                "de": h_phantram,
                                "ghi_chu": "",
                                "don_vi_kq":str(first_row.get("Đơn vị kết quả",""))
                            }
                            
                            for _, row in group.iterrows():
                                chi_tieu = str(row.get("Tên chỉ tiêu", ""))
                                c_do = str(row.get("C đo", "")).replace('.', ',')
                                kq_thuc = str(row.get("C thực", "")).replace('.', ',')
                                
                                ket_qua_list.append({
                                    "ngay": datetime.now().strftime("%d/%m/%Y"),
                                    "ten_mau": ten_mau_str,
                                    "chi_tieu": chi_tieu,
                                    "c_surr_truoc": c_surr_truoc,
                                    "c_surr_sau": c_surr_sau,
                                    "h_phantram": h_phantram,
                                    "v_khi": v_khi,
                                    "c_do": c_do,
                                    "kq_thuc": kq_thuc,
                                    "don_vi_kq":str(row.get("Đơn vị kết quả","")),
                                    "don_vi_cdo":str(row.get("Đơn vị C đo","")),
                                    "don_vi_csurr":str(row.get("Đơn vị Csurr",""))
                                })
                                
                                chi_tieu_upper = chi_tieu.upper()
                                slug = re.sub(r'\W+', '', chi_tieu_upper.lower())
                                mau_dict[f"cdo_{slug}"] = c_do
                                mau_dict[f"kq_{slug}"] = kq_thuc
                                
                                if "BENZENE" in chi_tieu_upper or "BENZEN" in chi_tieu_upper:
                                    mau_dict["cdo_benzen"] = c_do
                                    mau_dict["kq_benzen"] = kq_thuc
                                elif "TOLUENE" in chi_tieu_upper or "TOLUEN" in chi_tieu_upper:
                                    mau_dict["cdo_toluen"] = c_do
                                    mau_dict["kq_toluen"] = kq_thuc
                                    
                            danh_sach_mau.append(mau_dict)

                        if template_file.name.endswith('.docx'):
                            if st.button("🚀 Bơm Dữ Liệu Lập Biên Bản (Word)", type="primary"):
                                try:
                                    from docxtpl import DocxTemplate
                                    import io
                                    
                                    doc = DocxTemplate(template_file)
                                    doc.render({"danh_sach_mau": danh_sach_mau, "ket_qua": ket_qua_list})
                                    
                                    bio = io.BytesIO()
                                    doc.save(bio)
                                    bio.seek(0)
                                    
                                    st.success("🎉 Cập nhật số liệu thành công!")
                                    st.download_button("📥 Tải Xuống Biên Bản (.docx)", data=bio, file_name=f"Bien_Ban_{datetime.now().strftime('%Y%m%d_%H%M')}.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
                                except ImportError:
                                    st.error("⚠️ Máy chủ chưa cài thư viện 'docxtpl'. Hãy nhớ thêm 'docxtpl' vào file requirements.txt trên GitHub nhé!")
                                    
                        elif template_file.name.endswith('.xlsx'):
                            if st.button("🚀 Bơm Dữ Liệu Lập Biên Bản (Excel)", type="primary"):
                                import openpyxl
                                from copy import copy
                                import io
                                
                                wb = openpyxl.load_workbook(template_file)
                                ws = wb.active
                                
                                template_row_idx = None
                                template_cells = []
                                for r in range(1, ws.max_row + 1):
                                    for c in range(1, ws.max_column + 1):
                                        val = str(ws.cell(row=r, column=c).value or "")
                                        if "{{" in val:
                                            template_row_idx = r
                                            template_cells = [ws.cell(row=r, column=col).value for col in range(1, ws.max_column + 1)]
                                            break
                                    if template_row_idx:
                                        break
                                        
                                if template_row_idx:
                                    is_nuoc = any(isinstance(v, str) and ("kq_benzen" in v or "kq_toluen" in v or "mau.ky_hieu" in v) for v in template_cells)
                                    data_loop = danh_sach_mau if is_nuoc else ket_qua_list
                                    
                                    original_styles = []
                                    for col in range(1, ws.max_column + 1):
                                        cell_obj = ws.cell(row=template_row_idx, column=col)
                                        original_styles.append({
                                            "font": copy(cell_obj.font),
                                            "border": copy(cell_obj.border),
                                            "fill": copy(cell_obj.fill),
                                            "number_format": copy(cell_obj.number_format),
                                            "alignment": copy(cell_obj.alignment)
                                        })
                                    
                                    if len(data_loop) > 1 and ws.max_row > template_row_idx:
                                        ws.move_range(f"A{template_row_idx+1}:{ws.cell(ws.max_row, ws.max_column).coordinate}", rows=len(data_loop)-1, translate=True)
                                    
                                    current_row = template_row_idx
                                    for item in data_loop:
                                        for col_idx, cell_val in enumerate(template_cells, start=1):
                                            new_val = cell_val
                                            if cell_val and isinstance(cell_val, str):
                                                new_val = re.sub(r'\{%p.*?%\}', '', new_val).strip()
                                                new_val = re.sub(r'\{%.*?%\}', '', new_val).strip()
                                                
                                                matches = re.findall(r'\{\{\s*(?:row\.|mau\.)?(\w+)\s*\}\}', new_val)
                                                for m in matches:
                                                    replacement = str(item.get(m, ""))
                                                    new_val = re.sub(r'\{\{\s*(?:row\.|mau\.)?' + m + r'\s*\}\}', replacement, new_val)
                                                    
                                                if isinstance(new_val, str) and re.match(r'^-?\d+(?:,\d+)?$', new_val.strip()):
                                                    try:
                                                        new_val = float(new_val.strip().replace(',', '.'))
                                                    except: pass
                                                
                                                if new_val == "": new_val = None
                                                    
                                            new_cell = ws.cell(row=current_row, column=col_idx)
                                            new_cell.value = new_val
                                            
                                            if current_row >= template_row_idx:
                                                style = original_styles[col_idx - 1]
                                                new_cell.font = copy(style["font"])
                                                new_cell.border = copy(style["border"])
                                                new_cell.fill = copy(style["fill"])
                                                new_cell.number_format = copy(style["number_format"])
                                                new_cell.alignment = copy(style["alignment"])
                                                
                                        current_row += 1
                                                
                                    bio = io.BytesIO()
                                    wb.save(bio)
                                    bio.seek(0)
                                    
                                    st.success("🎉 Cập nhật số liệu thành công!")
                                    st.download_button("📥 Tải Xuống Biên Bản (.xlsx)", data=bio, file_name=f"Bien_Ban_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                                else:
                                    st.error("⚠️ Không tìm thấy vị trí chèn số liệu. Bạn hãy ghi đúng các thẻ (VD: `{{ row.ten_mau }}`) vào 1 dòng duy nhất trên Form Excel nhé.")
                except Exception as e:
                    st.error(f"❌ Có lỗi xảy ra trong quá trình xử lý: {e}")

    with tab_qr: 
        st.markdown("<div class='sub-title'>🏷️ Sinh Mã Vạch QR Tự Động</div>", unsafe_allow_html=True)
        st.write("Mô đun in tem dán mã vạch (Barcode/QR code) hàng loạt đang chờ tích hợp.")

elif menu == "🧪 Kiểm soát Hóa chất":
    st.markdown("<h1 class='main-title'>🧪 Kho hóa chất và Quản lý nhập liệu</h1>", unsafe_allow_html=True)
    if st.session_state.get('chem_error'): st.error(st.session_state.chem_error)
    base = st.session_state.df_chem.copy()
    
    # DASHBOARD THỐNG KÊ (MỚI)
    if not base.empty:
        col1, col2, col3, col4 = st.columns(4)
        col1.markdown(f'<div class="kpi" style="--accent:#239f8d; min-height:110px; padding:15px"><div class="kpi-label">Tổng số hóa chất</div><div class="kpi-number">{len(base)}</div></div>', unsafe_allow_html=True)
        col2.markdown(f'<div class="kpi" style="--accent:#5d8fbe; min-height:110px; padding:15px"><div class="kpi-label">Chuẩn phân tích</div><div class="kpi-number">{len(base[base["Phân Loại"] == "Chất chuẩn phân tích"])}</div></div>', unsafe_allow_html=True)
        col3.markdown(f'<div class="kpi" style="--accent:#8b82be; min-height:110px; padding:15px"><div class="kpi-label">Nội chuẩn/Surrogate</div><div class="kpi-number">{len(base[base["Phân Loại"] == "Chất chuẩn (IS/Surrogate)"])}</div></div>', unsafe_allow_html=True)
        col4.markdown(f'<div class="kpi" style="--accent:#dbab5b; min-height:110px; padding:15px"><div class="kpi-label">Cần rà soát</div><div class="kpi-number">{len(base[base["Cần Kiểm Tra"].astype(str).str.strip() != ""]) if "Cần Kiểm Tra" in base.columns else 0}</div></div>', unsafe_allow_html=True)
        st.write("")

    # NÂNG CẤP FORM THÊM THỦ CÔNG
    with st.expander('➕ Thêm hóa chất thủ công'):
        with st.form('stock_add_form'):
            col_add1, col_add2 = st.columns(2)
            with col_add1:
                name = st.text_input('Tên hóa chất *')
                lot = st.text_input('Số lô (Lot)')
                system = st.selectbox('Hệ máy', CHEM_SYSTEMS)
                kind = st.selectbox('Phân loại', ['Chất chuẩn phân tích'] + CHEM_TYPES)
            with col_add2:
                cas = st.text_input('CAS (Ví dụ: 59-50-7)')
                maker = st.text_input('Nhà sản xuất')
                conc = st.number_input('Nồng độ', min_value=0.0, value=0.0, format="%.2f")
                unit = st.selectbox('Đơn vị nồng độ', list(U_GROUPS['Nồng độ dung dịch']))
                
            if st.form_submit_button('Thêm vào kho', disabled=bool(st.session_state.get('chem_error'))):
                try:
                    if not name.strip(): raise ValueError('Cần nhập tên hóa chất.')
                    new_record = {
                        'Tên Hóa Chất': name.strip(), 'Số Lô (Lot)': lot, 'Hệ Máy': system, 
                        'Phân Loại': kind, 'Tình Trạng Kho': 'Chưa xác nhận lượng tồn',
                        'CAS': cas.strip(), 'Nhà Sản Xuất': maker.strip(), 
                        'Nồng Độ': conc if conc > 0 else None,
                        'Đơn Vị Nồng Độ': unit if conc > 0 else ''
                    }
                    row = stock_frame(pd.DataFrame([new_record]))
                    updated = pd.concat([base, row], ignore_index=True)
                    save_chemical_data(updated); st.session_state.df_chem = stock_frame(updated); st.rerun()
                except Exception as exc: st.error(str(exc))
                
    st.markdown("<div class='section-title'>📥 Nạp dữ liệu tự động (PDF / CSV)</div>", unsafe_allow_html=True)
    with st.container(border=True):
        col_up1, col_up2 = st.columns(2)
        with col_up1:
            upload_pdf = st.file_uploader('Tải lên PDF (Hóa chất chuẩn GC)', type=['pdf'], key='stock_pdf')
        with col_up2:
            upload_csv = st.file_uploader('Hoặc tải lên CSV (File bóc tách)', type=['csv'], key='stock_csv')

        parsed = None
        file_name_display = ""
        
        if upload_pdf:
            try:
                @st.cache_data(show_spinner=False, max_entries=3)
                def cached_stock(data, name): return read_stock_pdf(data, name)
                parsed = cached_stock(upload_pdf.getvalue(), upload_pdf.name)
                file_name_display = upload_pdf.name
            except Exception as exc: st.error(f"Lỗi đọc PDF: {exc}")
            
        elif upload_csv:
            try:
                parsed = pd.read_csv(upload_csv)
                parsed = stock_frame(parsed)
                file_name_display = upload_csv.name
            except Exception as exc: st.error(f"Lỗi đọc CSV: {exc}")

        # TIẾP NHẬN DỮ LIỆU TỪ FILE VÀ HIỂN THỊ TRƯỚC
        if parsed is not None and not parsed.empty:
            st.success(f'✔️ Đã đọc được {len(parsed)} dòng từ file {file_name_display}.')
            st.caption('Chọn dòng cần nhập vào kho chung. Chú ý các dòng có ghi chú trong cột Cần Kiểm Tra.')
            preview = parsed.copy(); preview.insert(0, 'Nhập', False)
            
            col_config = {
                'Hạn Sử Dụng': st.column_config.DateColumn(),
                'Nồng Độ': st.column_config.NumberColumn(),
                'Đơn Vị Nồng Độ': st.column_config.SelectboxColumn(options=list(U_GROUPS['Nồng độ dung dịch']))
            }
            disabled_cols = ['ID Nguồn', 'Nguồn PDF', 'Trang PDF', 'Dòng PDF', 'HSD Gốc', 'Tình Trạng Gốc', 'Quy Cách Gốc']
            
            edited_preview = st.data_editor(preview, num_rows='fixed', hide_index=True, use_container_width=True, 
                                    key='stock_preview_' + hashlib.sha256(str(file_name_display).encode()).hexdigest(),
                                    column_config=col_config, disabled=disabled_cols)
                                    
            if upload_pdf:
                st.download_button('📥 Tải kết quả bóc tách PDF (CSV)', parsed.to_csv(index=False).encode('utf-8-sig'), 'Kho_hoa_chat_tu_PDF.csv', 'text/csv')
                
            if st.button('🚀 Ghi các dòng đã chọn vào Kho chung', type="primary", disabled=bool(st.session_state.get('chem_error'))):
                try:
                    incoming = edited_preview[edited_preview['Nhập']].drop(columns='Nhập')
                    if incoming.empty: raise ValueError('Bạn chưa tick chọn dòng nào để nhập.')
                    updated = stock_merge(base, incoming)
                    save_chemical_data(updated)
                    st.session_state.df_chem = stock_frame(updated)
                    st.success("Nhập kho thành công!")
                    st.rerun()
                except Exception as exc: st.error(str(exc))
    
    # QUẢN LÝ KHO CHUNG & CẢNH BÁO
    if base.empty:
        st.info('Kho chưa có dữ liệu. Hãy thêm thủ công hoặc tải lên file PDF/CSV.')
    else:
        st.markdown("<div class='section-title' style='margin-top:20px'>📋 Danh sách Hóa chất & Vật tư</div>", unsafe_allow_html=True)
        expires = pd.to_datetime(base['Hạn Sử Dụng'], errors='coerce')
        today = pd.Timestamp.now().normalize()
        
        expired = base[expires < today]
        soon = base[(expires >= today) & (expires <= today + pd.Timedelta(days=30))]
        needs_check = base[base['Cần Kiểm Tra'].astype(str).str.strip() != ''] if 'Cần Kiểm Tra' in base.columns else pd.DataFrame()
        
        with st.expander('⚠️ Bảng Cảnh báo (Hạn sử dụng & Cần đối chiếu thông tin)'):
            if not expired.empty or not soon.empty:
                st.markdown("🔴 **Hóa chất Đã hết hạn / Sắp hết hạn:**")
                st.dataframe(pd.concat([expired, soon]), use_container_width=True)
            if not needs_check.empty:
                st.markdown("🔍 **Hóa chất cần kiểm tra lại thông tin (Sai CAS, lỗi ngày tháng...):**")
                st.dataframe(needs_check[['Tên Hóa Chất', 'Nhà Sản Xuất', 'CAS', 'Cần Kiểm Tra']], use_container_width=True)

        with st.container(border=True):
            col_f1, col_f2 = st.columns([2, 1])
            with col_f1: search = st.text_input('🔍 Lọc theo Tên / CAS / Nhà sản xuất')
            with col_f2: systems = st.multiselect('Lọc Hệ máy', sorted(set(base['Hệ Máy'].dropna().astype(str))))
            
            mask = pd.Series(True, index=base.index)
            if search: mask &= base[['Tên Hóa Chất', 'CAS', 'Nhà Sản Xuất']].fillna('').astype(str).apply(lambda c: c.str.contains(search, case=False, regex=False)).any(axis=1)
            if systems: mask &= base['Hệ Máy'].isin(systems)
            selected = base[mask].copy()
            
            st.caption('Sửa cột Đơn Vị Nồng Độ là sửa khai báo dữ liệu gốc. Muốn đổi đơn vị và giữ nguyên nồng độ thực, hãy dùng bảng Quy đổi bên dưới.')
            
            edited_inventory = st.data_editor(selected, num_rows='fixed', hide_index=True, use_container_width=True, key='stock_edit',
                column_config={
                    'Hạn Sử Dụng': st.column_config.DateColumn(),
                    'Ngày Mở Nắp': st.column_config.DateColumn(),
                    'Hệ Máy': st.column_config.SelectboxColumn(options=CHEM_SYSTEMS),
                    'Tình Trạng Kho': st.column_config.SelectboxColumn(options=CHEM_STATUS + ['Chưa xác nhận lượng tồn']),
                    'Đơn Vị Nồng Độ': st.column_config.SelectboxColumn(options=list(U_GROUPS['Nồng độ dung dịch']))
                }, disabled=['ID Nguồn', 'Nguồn PDF', 'Trang PDF', 'Dòng PDF'])
                
            if st.button('💾 Lưu các chỉnh sửa vào Kho', type="primary", disabled=bool(st.session_state.get('chem_error'))):
                try:
                    updated = base.copy(); updated.loc[selected.index, edited_inventory.columns] = edited_inventory.values
                    save_chemical_data(updated); st.session_state.df_chem = stock_frame(updated); st.rerun()
                except Exception as exc: st.error(str(exc))
                
            with st.expander('🔄 Công cụ Quy đổi nồng độ (Không ghi đè giá trị gốc)'):
                target = st.selectbox('Đơn vị xem nồng độ mong muốn', list(U_GROUPS['Nồng độ dung dịch']), index=1)
                view = selected[['Tên Hóa Chất', 'Nồng Độ', 'Đơn Vị Nồng Độ']].copy()
                def cv(row):
                    try: return u_convert(row['Nồng Độ'], row['Đơn Vị Nồng Độ'], target)
                    except (ValueError, TypeError): return None
                view['Nồng độ quy đổi'] = view.apply(cv, axis=1); view['Đơn vị đích'] = target
                st.dataframe(view, use_container_width=True)
                
            st.download_button('📥 Xuất toàn bộ Kho ra CSV', stock_serial(base).to_csv(index=False).encode('utf-8-sig'), 'Kho_hoa_chat_Hien_Tai.csv', 'text/csv')

elif menu == "🔄 Quy đổi đơn vị":
    unit_utility()

elif menu == "⚙️ Cấu hình Hệ thống":
    with st.expander('Quy đổi đơn vị thư viện MDL/LOQ để tải xuống'):
        target_limit=st.selectbox('Đơn vị đích cho các dòng tương thích',[u for g in U_GROUPS.values() for u in g],key='limit_target')
        if st.button('Tạo bản thư viện quy đổi'):
            converted=st.session_state.df_limit.copy()
            errors=[]
            for index,row in converted.iterrows():
                try:
                    factor=u_convert(1,row.get('Đơn Vị',''),target_limit)
                    values={}
                    for col in ['MDL','LOQ']:
                        val=row.get(col)
                        if pd.notna(val) and str(val).strip():values[col]=float(str(val).replace(',','.'))*factor
                    for col,value in values.items():converted.at[index,col]=value
                    converted.at[index,'Đơn Vị']=target_limit
                except (ValueError,TypeError):errors.append(index+1)
            st.dataframe(converted,use_container_width=True)
            if errors:st.info(f'{len(errors)} dòng khác đại lượng/thiếu đơn vị được giữ nguyên, không đổi nhãn.')
            st.download_button('Tải thư viện đã quy đổi CSV',converted.to_csv(index=False).encode('utf-8-sig'),'Thu_vien_quy_doi.csv','text/csv')
    st.markdown("<h1 class='main-title'>⚙️ Cấu hình & Quản trị Hệ thống</h1>", unsafe_allow_html=True)
    
    st.markdown("<div class='sub-title'>Quản lý Thư viện MDL & LOQ</div>", unsafe_allow_html=True)
    if st.session_state.get('results_stale'):
        st.warning("⚠️ Thư viện đã bị thay đổi! Vui lòng quay lại tab Vận hành máy đo và bấm 'Tính lại' để có kết quả mới nhất.")
        
    st.info("💡 Bạn có thể chỉnh sửa, thêm, xóa các mức giới hạn trực tiếp trên bảng. Nhớ ấn **Lưu thay đổi**.")
    
    with st.container(border=True):
        df_current_limit = st.session_state.df_limit.copy()
        if df_current_limit.empty:
            df_current_limit = pd.DataFrame(columns=["Nền Mẫu", "Tên Chất", "MDL", "LOQ", "Đơn Vị"])
            df_current_limit.loc[0] = ["", "", "", "", ""]
            
        edited_limit = st.data_editor(
            df_current_limit, 
            num_rows="dynamic", 
            use_container_width=True,
            key="limit_editor",
            height=350
        )
        
        if st.button("💾 Lưu thay đổi Thư viện", type="primary"):
            try:
                st.cache_data.clear() 
                edited_limit = edited_limit[edited_limit["Tên Chất"].str.strip() != ""] 
                
                conn.update(spreadsheet=SHEET_URL, worksheet="CauHinh_MDL_LOQ", data=edited_limit)
                st.session_state.df_limit = edited_limit
                if 'results' in st.session_state:
                    st.session_state.results_stale = True
                st.success("🎉 Đã lưu thư viện lên Google Sheets thành công!")
                st.rerun()
            except Exception as e:
                st.error(f"⚠️ Lỗi kết nối Google Sheets: {e}")

    with st.container(border=True):
        st.markdown("<div class='sub-title'>Cập nhật Hàng loạt (Import Excel)</div>", unsafe_allow_html=True)
        limit_file = st.file_uploader("Kéo thả file Bảng giới hạn (.xlsx)", type=["xlsx"])
        if limit_file:
            try:
                xls = pd.ExcelFile(limit_file)
                limit_data = []
                for sheet in xls.sheet_names:
                    df_sheet = pd.read_excel(xls, sheet_name=sheet, header=None)
                    header_idx = -1
                    c_ten, c_mdl, c_loq = None, None, None
                    
                    for r in range(min(20, len(df_sheet))):
                        row_vals = [str(val).lower() for val in df_sheet.iloc[r].values]
                        c_ten_temp = next((i for i, v in enumerate(row_vals) if 'tên' in v or 'hợp chất' in v), None)
                        c_mdl_temp = next((i for i, v in enumerate(row_vals) if 'mdl' in v), None)
                        c_loq_temp = next((i for i, v in enumerate(row_vals) if 'loq' in v), None)
                        
                        if c_ten_temp is not None and (c_mdl_temp is not None or c_loq_temp is not None):
                            header_idx, c_ten, c_mdl, c_loq = r, c_ten_temp, c_mdl_temp, c_loq_temp
                            break
                    
                    if header_idx != -1:
                        unit = "Chưa rõ"
                        if c_loq is not None:
                            unit_match = re.search(r'\((.*?)\)', str(df_sheet.iloc[header_idx, c_loq]))
                            if unit_match: unit = unit_match.group(1)
                        if unit == "Chưa rõ" and c_mdl is not None:
                            unit_match = re.search(r'\((.*?)\)', str(df_sheet.iloc[header_idx, c_mdl]))
                            if unit_match: unit = unit_match.group(1)

                        s_lower = sheet.lower()
                        nen_mau = "KT" if "thải" in s_lower else ("KXQ" if "xung quanh" in s_lower else ("KLV" if "làm việc" in s_lower else ("NS" if "nước" in s_lower or "voc" in s_lower else sheet)))
                        
                        for r in range(header_idx + 1, len(df_sheet)):
                            ten_val = df_sheet.iloc[r, c_ten]
                            if pd.isna(ten_val) or str(ten_val).strip() == "" or str(ten_val).lower() == 'nan': continue
                            
                            val_mdl = str(df_sheet.iloc[r, c_mdl]).replace(',', '.').strip() if c_mdl is not None and pd.notna(df_sheet.iloc[r, c_mdl]) else ""
                            val_loq = str(df_sheet.iloc[r, c_loq]).replace(',', '.').strip() if c_loq is not None and pd.notna(df_sheet.iloc[r, c_loq]) else ""
                            if val_mdl.lower() == 'nan': val_mdl = ""
                            if val_loq.lower() == 'nan': val_loq = ""
                            
                            if val_mdl or val_loq:
                                limit_data.append({"Nền Mẫu": nen_mau, "Tên Chất": str(ten_val).strip(), "MDL": val_mdl, "LOQ": val_loq, "Đơn Vị": unit})
                
                if limit_data:
                    df_limit_new = pd.DataFrame(limit_data)
                    st.success(f"✔️ Đã nhận diện {len(df_limit_new)} chỉ tiêu mới.")
                    st.dataframe(df_limit_new, use_container_width=True)
                    
                    if st.button("🚀 Trộn và Lưu vào Thư viện chung", type="primary"):
                        try:
                            st.cache_data.clear() 
                            combined_df = pd.concat([st.session_state.df_limit, df_limit_new], ignore_index=True)
                            combined_df = combined_df.drop_duplicates(subset=['Nền Mẫu', 'Tên Chất'], keep='last').reset_index(drop=True)
                            
                            conn.update(spreadsheet=SHEET_URL, worksheet="CauHinh_MDL_LOQ", data=combined_df)
                            st.session_state.df_limit = combined_df
                            if 'results' in st.session_state:
                                st.session_state.results_stale = True
                            st.success(f"🎉 Ghi đè thành công! Tổng bộ nhớ: {len(combined_df)} chỉ tiêu.")
                            st.rerun()
                        except Exception as sheet_err:
                            st.error(f"⚠️ Lỗi kết nối: {sheet_err}")
                else: st.error("Không tìm thấy cấu trúc bảng hợp lệ (Cột Tên / Cột MDL / Cột LOQ).")
            except Exception as e: st.error(f"Lỗi đọc file: {e}")
